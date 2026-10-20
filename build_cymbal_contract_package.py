"""
build_cymbal_contract_package.py

Generates the full-length, multi-subsection Cymbal Aerospace Corporation contract package
in `specs/04_golden_fixtures/`:
1. `Cymbal_GP1_Standard.docx` (and `Boeing_GP1_Standard.docx` for path compatibility)
   - Formal Cover Header, Master Subcontract Metadata Table, Recitals (WHEREAS...),
     Part I (GP1), Part II (SP3), Part III (SP5) Section Dividers, 12 full-length multi-subsection
     aerospace legal clauses (`(a)`, `(b)`, `(c)`), and Bilateral Execution Signature Block.
2. `Vendor_Redlined_GP1.docx` (and `Cymbal_Vendor_Redlined_GP1.docx`)
   - Full-length 12-clause contract with surgical native Word OOXML Track Changes
     (`<w:del>`, `<w:ins w:author="Apex AeroSystems Legal">`) and clickable `<w:comment>`
     margin bubbles across the 4 disputed clauses (`GP1-12`, `GP1-19`, `SP3-04`, `SP5-03`).
3. `Cymbal_SOW_Schedule_A_Pricing.docx`
   - Companion Statement of Work (`SOW-AV-8921`) and Schedule A Firm-Fixed-Price Unit Pricing
     Table ($45,000,000 ceiling across 300 shipsets of FCC-900 Flight Control Computers,
     containing the explicit §3.1 FFP price lock referenced by the SP3-04 vs GP1-03 conflict).
4. `Apex_AeroSystems_Redline_Transmittal_Memo.docx`
   - Formal Vendor Exception & Redline Transmittal Memorandum from Apex AeroSystems' General Counsel
     to Cymbal Aerospace Procurement summarizing their 4 commercial/regulatory exceptions.
5. Updates `specs/04_golden_fixtures/expected_extracted_diffs.json` with the exact output of
   `parse_redlined_docx("specs/04_golden_fixtures/Vendor_Redlined_GP1.docx")`.
"""
from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
import shutil
from typing import Any, Dict, List, Tuple
import zipfile

from docx import Document
from lxml import etree

from boeing_clm.ingestion.docx_track_changes import parse_redlined_docx

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
COMMENTS_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
COMMENTS_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"

VENDOR_AUTHOR = "Apex AeroSystems Legal"
REDLINE_DATE = "2026-10-15T14:30:00Z"

REPO_ROOT = Path(__file__).resolve().parent
GOLDEN_DIR = REPO_ROOT / "specs" / "04_golden_fixtures"

SHAREPOINT_SITE_URL = (
    "https://36zrbd.sharepoint.com/sites/CymbalAerospace/Shared%20Documents/BA-2026-AV-8921"
)


# ---------------------------------------------------------------------------
# Full-Length Multi-Subsection Clause Definitions for Cymbal Aerospace
# Each clause has:
#   - header: Heading 2 metadata string parsed by `parse_heading_metadata`
#   - part_heading: Optional Heading 1 divider placed before the clause
#   - segments: List of ("same" | "del" | "ins", text) tuples
#   - comment: Optional vendor margin comment injected into word/comments.xml
#   - fallback_tier_1_text: Tier-1 compromise text for BigQuery proforma_clauses
#   - fallback_tier_2_text: Tier-2 compromise text for BigQuery proforma_clauses
#   - regulatory_lockout: bool
#   - regulatory_citation: str
#   - guardrail_summary: str
# ---------------------------------------------------------------------------
CLAUSES_SPEC: List[Dict[str, Any]] = [
    {
        "clause_id": "GP1-01",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §1.1",
        "clause_title": "Definitions & Order of Precedence",
        "sme_team": "PROCUREMENT",
        "sme_owner_label": "Procurement & Subcontracts",
        "risk_level": "NONE",
        "regulatory_flag": "Cymbal Corporate GP1 Standard",
        "part_heading": "PART I — GENERAL PROVISIONS (CYMBAL AEROSPACE GP1)",
        "segments": [
            (
                "same",
                "1.1 (a) Definitions: \"Buyer\" means Cymbal Aerospace Corporation or its wholly-owned aerospace subsidiary identified on the face of the Master Purchase Agreement (`BA-2026-AV-8921`). \"Seller\" means Apex AeroSystems Inc. or the legal entity furnishing Goods and flight-critical software hereunder. \"Goods\" means the Triple-Redundant Fly-By-Wire Flight Control Computers (`FCC-900`), Line Replaceable Units (`LRUs`), embedded DAL-A flight control firmware, and associated technical deliverables specified in Statement of Work (`SOW-AV-8921`) Schedule A. (b) Order of Precedence: In the event of any ambiguity or inconsistency among Contract documents, Special Provisions (`SP3` International & Tariff Provisions and `SP5` Cybersecurity & Defense Provisions) shall take precedence over these General Provisions (`GP1`), expressly excluding the Firm-Fixed-Price (`FFP`) unit price locks and $45,000,000 aggregate ceiling set forth in GP1 §3.1 and SOW Schedule A, which shall govern over any conflicting Special Provision unless modified by a formal bilateral Change Order.",
            )
        ],
        "comment": None,
        "regulatory_lockout": False,
        "regulatory_citation": "Cymbal Corporate GP1 Standard",
        "guardrail_summary": "Maintain GP1 §3.1 Firm-Fixed-Price precedence over Special Provisions.",
    },
    {
        "clause_id": "GP1-03",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §3.1",
        "clause_title": "Firm-Fixed-Price SOW Cap & Pricing Integrity",
        "sme_team": "PROCUREMENT",
        "sme_owner_label": "Procurement & Cost Management",
        "risk_level": "NONE",
        "regulatory_flag": "Firm-Fixed-Price SOW Lock",
        "part_heading": None,
        "segments": [
            (
                "same",
                "3.1 (a) Firm-Fixed-Price Lock: Unless expressly modified by a bilateral Change Order executed by Buyer's Vice President of Avionics Procurement referencing this GP1 §3.1 and Statement of Work (`SOW-AV-8921`) Schedule A, all unit prices ($150,000 per shipset across 300 C780/C730-MAX shipsets) and the forty-five million dollar ($45,000,000.00) total Contract ceiling are strictly Firm-Fixed-Price (`FFP`) and shall not be subject to uncapped surcharges, tariff pass-throughs, customs duty adjustments, foreign exchange fluctuations, or raw material commodity indexing. (b) Invoicing & Payment Integrity: Seller shall submit itemized electronic invoices upon DDP acceptance of each production lot at Buyer's Everett or Renton final assembly line, payable Net Sixty (60) days. Any unilateral invoice surcharge not authorized by an executed GP1 §3.1 Change Order shall be automatically rejected.",
            )
        ],
        "comment": None,
        "regulatory_lockout": False,
        "regulatory_citation": "Firm-Fixed-Price SOW Lock",
        "guardrail_summary": "Strict Firm-Fixed-Price SOW lock ($45M ceiling). Any Special Provision tariff surcharge adjustment must be strictly capped at 2.0% aggregate adjustment or executed via bilateral Change Order.",
    },
    {
        "clause_id": "GP1-05",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §5.2",
        "clause_title": "Packing, Marking & Aerospace Preservation",
        "sme_team": "ENGINEERING",
        "sme_owner_label": "Supplier Logistics & Packaging",
        "risk_level": "NONE",
        "regulatory_flag": "ATA Spec 300 / D6-81628",
        "part_heading": None,
        "segments": [
            (
                "same",
                "5.2 (a) Flight-Critical Avionics Preservation: Seller shall clean, preserve, pack, mark, and ship all `FCC-900` Flight Control Computer LRUs in strict accordance with Cymbal Aerospace Supplier Packaging Specification `D6-81628` and Air Transport Association (`ATA`) Spec 300 Category I reusable container and MIL-STD-1686 Class 1 electro-static discharge (`ESD`) protection standards. (b) Traceability & Barcoding: Each individual LRU chassis and exterior shipping container shall bear permanent UID/GS1-128 serialized barcoding, shock-watch and tilt-indicator telemetry seals, desiccant humidity indicators, and an enclosed FAA Form 8130-3 Authorized Release Certificate. Any shipment arriving with tripped shock indicators or breached ESD shielding shall be rejected at Seller's sole expense.",
            )
        ],
        "comment": None,
        "regulatory_lockout": False,
        "regulatory_citation": "ATA Spec 300 / D6-81628",
        "guardrail_summary": "Require ATA Spec 300 Class 1 ESD protection and Cymbal D6-81628 packaging compliance.",
    },
    {
        "clause_id": "GP1-07",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §7.1",
        "clause_title": "Quality Management System (AS9100 Rev D)",
        "sme_team": "ENGINEERING",
        "sme_owner_label": "Supplier Quality Assurance",
        "risk_level": "NONE",
        "regulatory_flag": "AS9100 Rev D / D6-82479",
        "part_heading": None,
        "segments": [
            (
                "same",
                "7.1 (a) Accredited Aerospace QMS: Seller shall establish and maintain an IAQG-accredited `AS9100 Rev D` Quality Management System, `RTCA DO-178C` Design Assurance Level A (`DAL-A`) software lifecycle compliance, `RTCA DO-254` complex electronic hardware assurance, and Cymbal Aerospace Quality Management System Requirements for Suppliers (`D6-82479`) throughout Contract performance. (b) First Article Inspection & Sub-Tier Flowdown: Prior to initial production lot delivery, Seller shall complete a full `AS9102` First Article Inspection Report (`FAIR`) witnessed by Buyer's Supplier Quality Representative and shall flow down identical AS9100 Rev D and counterfeit electronic part avoidance (`AS5553`) controls to all sub-tier wafer foundries and forging sub-contractors.",
            )
        ],
        "comment": None,
        "regulatory_lockout": False,
        "regulatory_citation": "AS9100 Rev D / D6-82479",
        "guardrail_summary": "Mandatory AS9100 Rev D certification and Cymbal D6-82479 supplier quality flowdown.",
    },
    {
        "clause_id": "GP1-09",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §9.1",
        "clause_title": "Changes & Equitable Adjustment",
        "sme_team": "PROCUREMENT",
        "sme_owner_label": "Procurement & Subcontracts",
        "risk_level": "NONE",
        "regulatory_flag": "Cymbal Corporate GP1 Standard",
        "part_heading": None,
        "segments": [
            (
                "same",
                "9.1 (a) Directed Engineering & Schedule Changes: Buyer's Authorized Procurement Representative may at any time, by written Change Order without notice to sureties, direct changes within the general scope of this Contract in (i) drawings, avionics bus interface control documents (`ICDs`), or technical specifications, (ii) method of shipment or aerospace preservation, or (iii) place and schedule of delivery. Seller shall proceed immediately with performance of the Contract as changed. (b) Equitable Adjustment Claims: If any directed change under §9.1(a) causes a verified increase or decrease in the cost of or time required for performance, Seller must assert any claim for equitable adjustment in writing within thirty (30) calendar days from receipt of the written Change Order, supported by certified cost and pricing data.",
            )
        ],
        "comment": None,
        "regulatory_lockout": False,
        "regulatory_citation": "Cymbal Corporate GP1 Standard",
        "guardrail_summary": "Unilateral directed changes within scope require Authorized Procurement Representative written order.",
    },
    {
        "clause_id": "GP1-12",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §12.2",
        "clause_title": "Warranty & FAA Source Inspection",
        "sme_team": "ENGINEERING",
        "sme_owner_label": "Engineering & QA (FAA DER)",
        "risk_level": "HIGH",
        "regulatory_flag": "FAA 14 CFR Part 21 Lockout",
        "part_heading": None,
        "segments": [
            (
                "same",
                "12.2 (a) Airworthiness & Hardware/Firmware Warranty: Seller warrants all Goods, including `FCC-900` flight control hardware, triplex voting channels, and embedded `DO-178C DAL-A` avionics firmware, to be free from defects in design, material, and workmanship for ",
            ),
            (
                "del",
                "thirty-six (36) months from initial aircraft delivery to Cymbal Aerospace's airline customer",
            ),
            (
                "ins",
                "twelve (12) months from factory shipment",
            ),
            (
                "same",
                ", and Seller shall repair or replace defective flight-critical LRUs on an Aircraft-on-Ground (`AOG`) 24-hour turnaround basis at Seller's expense. (b) Regulatory Source Surveillance: Cymbal Aerospace and Federal Aviation Administration (`FAA`) representatives may conduct ",
            ),
            (
                "del",
                "unannounced source inspections at Seller's facility at any time",
            ),
            (
                "ins",
                "scheduled inspections upon thirty (30) days prior written notice",
            ),
            (
                "same",
                " to verify `AS9100 Rev D` quality compliance, `FAA 14 CFR Part 21.140` production certificate airworthiness conformity, and sub-tier test bench calibration records.",
            ),
        ],
        "comment": (
            "Apex AeroSystems Legal: Our standard commercial avionics warranty is 12 months from ex-works factory shipment, "
            "and unannounced facility audits disrupt Class-100 cleanroom assembly shifts; we require 30 days advance written notice."
        ),
        "fallback_tier_1_text": (
            "12.2 (a) Airworthiness & Hardware/Firmware Warranty: Seller warrants all Goods, including `FCC-900` flight control hardware, triplex voting channels, and embedded `DO-178C DAL-A` avionics firmware, to be free from defects in design, material, and workmanship for twenty-four (24) months from initial aircraft delivery to Cymbal Aerospace's airline customer or thirty (30) months from factory shipment, whichever occurs first, and Seller shall repair or replace defective flight-critical LRUs on an Aircraft-on-Ground (`AOG`) 24-hour turnaround basis at Seller's expense. (b) Regulatory Source Surveillance: Cymbal Aerospace and Federal Aviation Administration (`FAA`) representatives may conduct unannounced source inspections at Seller's facility at any time to verify `AS9100 Rev D` quality compliance, `FAA 14 CFR Part 21.140` production certificate airworthiness conformity, and sub-tier test bench calibration records."
        ),
        "fallback_tier_2_text": (
            "12.2 (a) Airworthiness & Hardware/Firmware Warranty: Seller warrants all Goods against defects in material, workmanship, and avionics firmware for twenty-four (24) months from factory shipment. (b) Regulatory Source Surveillance: Cymbal Aerospace and FAA representatives retain unannounced source inspection access at Seller's facility at any time under FAA 14 CFR Part 21.140, with 48-hour courtesy coordination for non-safety routine quality reviews."
        ),
        "regulatory_lockout": True,
        "regulatory_citation": "FAA 14 CFR Part 21 Lockout",
        "guardrail_summary": (
            "HARD REGULATORY LOCKOUT: FAA 14 CFR Part 21.140 and AS9100 Rev D strictly prohibit prior-notice restrictions on FAA/Cymbal Aerospace airworthiness source inspections. Warranty period minimum is 24 months from delivery (36 months baseline)."
        ),
    },
    {
        "clause_id": "GP1-15",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §15.1",
        "clause_title": "Termination for Convenience",
        "sme_team": "LEGAL",
        "sme_owner_label": "Corporate Legal & Commercial",
        "risk_level": "NONE",
        "regulatory_flag": "FAR 52.249-2 Aligned",
        "part_heading": None,
        "segments": [
            (
                "same",
                "15.1 (a) Unilateral Right to Terminate: Buyer may terminate all or any severable portion of this Contract for convenience at any time by written notice of termination to Seller. Upon receipt of such notice, Seller shall immediately stop all terminated work, place no further sub-tier purchase orders for raw materials or semiconductors, and protect all property in which Buyer has or may acquire an interest. (b) Termination Settlement & Cap: Seller's sole recovery upon termination for convenience, aligned with `FAR 52.249-2` principles, shall be limited to documented actual costs incurred for work-in-process prior to termination notice plus a reasonable profit on completed work, never to exceed the terminated portion of the GP1 §3.1 Firm-Fixed-Price Contract value, with zero recovery for anticipatory or consequential profits.",
            )
        ],
        "comment": None,
        "regulatory_lockout": False,
        "regulatory_citation": "FAR 52.249-2 Aligned",
        "guardrail_summary": "Preserve Buyer unilateral termination for convenience right aligned with FAR 52.249-2.",
    },
    {
        "clause_id": "GP1-19",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §19 & §24",
        "clause_title": "Limitation of Liability & IP Indemnity",
        "sme_team": "LEGAL",
        "sme_owner_label": "Corporate Legal & IP",
        "risk_level": "CRITICAL",
        "regulatory_flag": "Corporate Legal Playbook Tier-1",
        "part_heading": None,
        "segments": [
            (
                "same",
                "19.1 (a) Direct, Line-Stoppage & Airworthiness Recall Liability: Seller's aggregate liability to Cymbal Aerospace for direct damages, Everett/Renton final assembly line-stoppage costs, Fleet Airworthiness Directive (`AD`) grounding/retrofit campaigns, and airworthiness recall costs shall ",
            ),
            (
                "del",
                "be uncapped",
            ),
            (
                "ins",
                "not exceed one hundred percent (100%) of the total Purchase Order value ($45M)",
            ),
            (
                "same",
                ". 24.1 (b) Intellectual Property Indemnification: Seller shall defend, indemnify, and hold harmless Cymbal Aerospace and its airline customers against all third-party patent, copyright, trade secret, and intellectual property infringement claims arising from the manufacture, sale, or use of the Goods",
            ),
            (
                "del",
                " including system-level avionics bus integration",
            ),
            (
                "ins",
                " except where such claim arises from combination of Goods with any Cymbal Aerospace or third-party flight control bus",
            ),
            (
                "same",
                ", and if any `FCC-900` unit is enjoined, Seller shall at its sole expense procure a perpetual royalty-free license or re-architect the infringing avionics module to maintain FAA type-design airworthiness.",
            ),
        ],
        "comment": (
            "Apex AeroSystems Legal: Board risk policy requires a strict 1x PO liability cap ($45M) across all direct/recall damages "
            "and a standard commercial carveout excluding IP claims arising from combination of our LRUs with Cymbal Aerospace's flight control bus."
        ),
        "fallback_tier_1_text": (
            "19.1 (a) Direct, Line-Stoppage & Airworthiness Recall Liability: Seller's liability for general direct damages and final assembly line-stoppage shall be capped at two hundred percent (200%) of the total Purchase Order value ($90M), provided that liability for Fleet Airworthiness Directive (`AD`) recall, gross negligence, willful misconduct, and §24 IP indemnification shall remain strictly uncapped. 24.1 (b) Intellectual Property Indemnification: Seller shall defend, indemnify, and hold harmless Cymbal Aerospace and its airline customers against all third-party intellectual property infringement claims including system-level avionics bus integration, except where such claim arises solely from Cymbal Aerospace's proprietary flight control bus specifications."
        ),
        "fallback_tier_2_text": (
            "19.1 (a) Direct, Line-Stoppage & Airworthiness Recall Liability: Seller's liability for general direct damages shall be capped at one hundred fifty percent (150%) of the total Purchase Order value ($67.5M), with uncapped liability for airworthiness directives and IP indemnity. 24.1 (b) Intellectual Property Indemnification: Seller shall defend and indemnify Cymbal Aerospace against all third-party intellectual property infringement claims except where arising solely from Cymbal Aerospace-furnished specifications."
        ),
        "regulatory_lockout": False,
        "regulatory_citation": "Corporate Legal Playbook Tier-1",
        "guardrail_summary": (
            "Reject 1x PO ($45M) general cap. Playbook Tier-1 Compromise permits a 2x PO ($90M) super-cap on general direct damages only if airworthiness recall, gross negligence, and IP indemnity remain uncapped, and narrows IP carveout to claims arising solely from Cymbal Aerospace proprietary bus specifications."
        ),
    },
    {
        "clause_id": "GP1-28",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §28.1",
        "clause_title": "Governing Law & Dispute Resolution",
        "sme_team": "LEGAL",
        "sme_owner_label": "Corporate Legal",
        "risk_level": "NONE",
        "regulatory_flag": "Delaware / Washington Law Standard",
        "part_heading": None,
        "segments": [
            (
                "same",
                "28.1 (a) Governing Law & Exclusive Venue: This Contract and all claims, controversies, or disputes arising hereunder shall be governed by and construed in accordance with the substantive laws of the State of Delaware, excluding its conflict of laws rules and excluding the United Nations Convention on Contracts for the International Sale of Goods (`CISG`). Exclusive judicial venue shall lie in the state or federal courts located in Wilmington, Delaware or King County, Washington. (b) Continuity of Flight-Critical Supply: Pending final resolution of any dispute, claim, or Change Order valuation, Seller shall proceed diligently with uninterrupted manufacture, testing, and delivery of all `FCC-900` flight control units in accordance with Buyer's production schedule.",
            )
        ],
        "comment": None,
        "regulatory_lockout": False,
        "regulatory_citation": "Delaware / Washington Law Standard",
        "guardrail_summary": "Maintain Delaware or Washington state governing law and uninterrupted delivery during disputes.",
    },
    {
        "clause_id": "SP3-02",
        "proforma_doc": "SP3",
        "clause_ref": "SP3 §2.1",
        "clause_title": "Export Compliance (ITAR & EAR)",
        "sme_team": "TAX",
        "sme_owner_label": "Global Trade Controls",
        "risk_level": "NONE",
        "regulatory_flag": "22 CFR 120-130 (ITAR) / 15 CFR 730-774 (EAR)",
        "part_heading": "PART II — INTERNATIONAL & GLOBAL TRADE SPECIAL PROVISIONS (SP3)",
        "segments": [
            (
                "same",
                "2.1 (a) Export Control Regulatory Compliance: Seller shall comply with all applicable U.S. export control laws and regulations, including the International Traffic in Arms Regulations (`ITAR`, 22 CFR Parts 120–130) and the Export Administration Regulations (`EAR`, 15 CFR Parts 730–774), covering dual-use flight control algorithms (`ECCN 7A994` / `7E994`) and defense technical data. (b) Export Classification & Deemed Exports: Seller shall maintain active DDTC registration where applicable, furnish accurate ECCN/USML Commodity Jurisdiction classifications prior to first shipment, and prevent unauthorized deemed exports of Cymbal Aerospace technical data to foreign nationals without prior U.S. Department of State or Bureau of Industry and Security (`BIS`) authorization.",
            )
        ],
        "comment": None,
        "regulatory_lockout": False,
        "regulatory_citation": "22 CFR 120-130 (ITAR) / 15 CFR 730-774 (EAR)",
        "guardrail_summary": "Strict compliance with ITAR (22 CFR 120-130) and EAR (15 CFR 730-774).",
    },
    {
        "clause_id": "SP3-04",
        "proforma_doc": "SP3",
        "clause_ref": "SP3 §4.1",
        "clause_title": "Customs, Duties & International Tariffs",
        "sme_team": "TAX",
        "sme_owner_label": "Global Tax & Trade",
        "risk_level": "MOD",
        "regulatory_flag": "Cross-Clause SOW §3.1 Dependency",
        "part_heading": None,
        "segments": [
            (
                "same",
                "4.1 (a) Importer of Record & Customs Clearance: Seller shall act as Importer of Record for all imported sub-tier components, raw aerospace titanium forgings, and radiation-hardened semiconductor wafers. (b) Tariff & Duty Allocation: All unit prices in Statement of Work (`SOW-AV-8921`) Schedule A are ",
            ),
            (
                "del",
                "DDP (Incoterms 2020) Cymbal Aerospace Everett/Renton facility and inclusive of all import duties, customs fees, and tariffs",
            ),
            (
                "ins",
                "exclusive of any new or increased Section 232/301 tariffs on titanium forgings or semiconductor wafers exceeding five percent (5%), which shall be passed through 100% to Cymbal Aerospace",
            ),
            (
                "same",
                ", and Seller shall cooperate with Buyer's Global Trade team on U.S. CBP duty drawback documentation.",
            ),
        ],
        "comment": (
            "Apex AeroSystems Tax: Due to volatile Section 232/301 trade duties on European titanium forgings and Taiwanese semiconductor wafers, "
            "Seller cannot absorb open-ended tariff spikes under DDP pricing and requires 100% pass-through of tariff increases above 5%."
        ),
        "fallback_tier_1_text": (
            "4.1 (a) Importer of Record & Customs Clearance: Seller shall act as Importer of Record for all imported sub-tier components, raw aerospace titanium forgings, and radiation-hardened semiconductor wafers. (b) Tariff & Duty Allocation: All unit prices in Statement of Work (`SOW-AV-8921`) Schedule A are DDP (Incoterms 2020) Cymbal Aerospace Everett/Renton facility. Any verified new Section 232/301 tariff increase exceeding 5.0% on raw titanium forgings or semiconductor wafers shall be shared 50/50 between Cymbal Aerospace and Seller, strictly capped at a maximum aggregate adjustment of 2.0% of the GP1 §3.1 Firm-Fixed-Price SOW unit price."
        ),
        "fallback_tier_2_text": (
            "4.1 (a) Importer of Record & Customs Clearance: Seller shall act as Importer of Record for all imported sub-tier components. (b) Tariff & Duty Allocation: All unit prices are DDP Cymbal Aerospace Everett/Renton facility. Verified Section 232/301 tariff increases above 5.0% on raw titanium forgings may be shared 50/50 up to a maximum cap of 2.0% of the GP1 §3.1 Firm-Fixed-Price SOW unit price upon documentary customs proof."
        ),
        "regulatory_lockout": False,
        "regulatory_citation": "Cross-Clause SOW §3.1 Dependency",
        "guardrail_summary": (
            "CROSS-CLAUSE GUARDRAIL: Any tariff sharing or pass-through in SP3 §4.1 MUST be explicitly capped at a maximum aggregate adjustment of 2.0% to preserve the GP1 §3.1 Firm-Fixed-Price SOW Cap ($45M ceiling). Uncapped >5% pass-through triggers a CRITICAL cross-clause conflict with Procurement."
        ),
    },
    {
        "clause_id": "SP5-03",
        "proforma_doc": "SP5",
        "clause_ref": "SP5 §3.2",
        "clause_title": "Cybersecurity Incident Reporting (DFARS)",
        "sme_team": "CYBER",
        "sme_owner_label": "Supplier Cyber Defense",
        "risk_level": "HIGH",
        "regulatory_flag": "DFARS 252.204-7012 Lockout",
        "part_heading": "PART III — CYBERSECURITY & CONTROLLED UNCLASSIFIED INFORMATION SPECIAL PROVISIONS (SP5)",
        "segments": [
            (
                "same",
                "3.2 (a) CUI Safeguards & NIST SP 800-171 Rev 2 Compliance: Seller shall maintain an accredited Cybersecurity Maturity Model Certification (`CMMC 2.0 Level 2`) posture and implement all 110 security controls in `NIST SP 800-171 Rev 2` across all engineering networks storing or processing Cymbal Aerospace Covered Defense Information (`CDI`), Controlled Unclassified Information (`CUI`), or `FCC-900` flight control source code. (b) Mandatory Cyber Incident Reporting & Forensics: Pursuant to `DoD DFARS 252.204-7012`, Seller shall report any cybersecurity incident affecting Cymbal Aerospace CDI or technical data to Buyer's Cyber Security Operations Center (`CSOC`) and the DoD DIBNet portal within ",
            ),
            (
                "del",
                "seventy-two (72) hours of discovery",
            ),
            (
                "ins",
                "ten (10) business days following confirmation of a material breach",
            ),
            (
                "same",
                ", preserve affected system media for ninety (90) days, and provide raw forensic log images and malware telemetry upon request.",
            ),
        ],
        "comment": (
            "Apex AeroSystems CISO: Our internal Security Operations Center (SOC) requires 10 business days to complete forensic root-cause "
            "triage and confirm a material breach before issuing external customer notifications."
        ),
        "fallback_tier_1_text": (
            "3.2 (a) CUI Safeguards & NIST SP 800-171 Rev 2 Compliance: Seller shall maintain an accredited Cybersecurity Maturity Model Certification (`CMMC 2.0 Level 2`) posture and implement all 110 security controls in `NIST SP 800-171 Rev 2` across all engineering networks storing or processing Cymbal Aerospace Covered Defense Information (`CDI`), Controlled Unclassified Information (`CUI`), or `FCC-900` flight control source code. (b) Mandatory Cyber Incident Reporting & Forensics: Pursuant to `DoD DFARS 252.204-7012`, Seller shall report any cybersecurity incident affecting Cymbal Aerospace CDI or technical data within seventy-two (72) hours of discovery via preliminary notification, followed by root-cause forensic log images within five (5) business days."
        ),
        "fallback_tier_2_text": (
            "3.2 (a) CUI Safeguards & NIST SP 800-171 Rev 2 Compliance: Seller shall maintain NIST SP 800-171 controls across all systems processing Cymbal Aerospace CDI/CUI. (b) Mandatory Cyber Incident Reporting & Forensics: Seller shall submit mandatory DFARS 252.204-7012 cyber incident notification within seventy-two (72) hours of discovery, with full forensic packet delivery within seven (7) business days."
        ),
        "regulatory_lockout": True,
        "regulatory_citation": "DFARS 252.204-7012 Lockout",
        "guardrail_summary": (
            "HARD REGULATORY LOCKOUT: DoD DFARS 252.204-7012 mandates cyber incident reporting within 72 hours of discovery. Extending initial notification to 10 business days is a federal defense flowdown violation."
        ),
    },
]


def _append_run(p_el: etree._Element, text: str) -> None:
    r = etree.SubElement(p_el, f"{{{W_NS}}}r")
    t = etree.SubElement(r, f"{{{W_NS}}}t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text


def _append_del(p_el: etree._Element, rev_id: int, text: str) -> None:
    del_el = etree.SubElement(
        p_el,
        f"{{{W_NS}}}del",
        {
            f"{{{W_NS}}}id": str(rev_id),
            f"{{{W_NS}}}author": VENDOR_AUTHOR,
            f"{{{W_NS}}}date": REDLINE_DATE,
        },
    )
    r = etree.SubElement(del_el, f"{{{W_NS}}}r")
    dt = etree.SubElement(r, f"{{{W_NS}}}delText")
    dt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    dt.text = text


def _append_ins(p_el: etree._Element, rev_id: int, text: str) -> None:
    ins_el = etree.SubElement(
        p_el,
        f"{{{W_NS}}}ins",
        {
            f"{{{W_NS}}}id": str(rev_id),
            f"{{{W_NS}}}author": VENDOR_AUTHOR,
            f"{{{W_NS}}}date": REDLINE_DATE,
        },
    )
    r = etree.SubElement(ins_el, f"{{{W_NS}}}r")
    t = etree.SubElement(r, f"{{{W_NS}}}t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text


def _inject_comments_part(docx_bytes: bytes, comments_list: List[Dict[str, Any]]) -> bytes:
    """Inject word/comments.xml and package relationships into a .docx zip archive."""
    in_mem = BytesIO(docx_bytes)
    out_mem = BytesIO()

    nsmap = {"w": W_NS}
    comments_root = etree.Element(f"{{{W_NS}}}comments", nsmap=nsmap)
    for c in comments_list:
        c_el = etree.SubElement(
            comments_root,
            f"{{{W_NS}}}comment",
            {
                f"{{{W_NS}}}id": str(c["id"]),
                f"{{{W_NS}}}author": c["author"],
                f"{{{W_NS}}}date": c["date"],
            },
        )
        p_el = etree.SubElement(c_el, f"{{{W_NS}}}p")
        _append_run(p_el, c["text"])

    comments_xml_bytes = etree.tostring(
        comments_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    with zipfile.ZipFile(in_mem, "r") as zin, zipfile.ZipFile(out_mem, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                ct_root = etree.fromstring(data)
                exists = any(
                    el.get("PartName") == "/word/comments.xml"
                    for el in ct_root.findall(f"{{{CT_NS}}}Override")
                )
                if not exists:
                    etree.SubElement(
                        ct_root,
                        f"{{{CT_NS}}}Override",
                        {
                            "PartName": "/word/comments.xml",
                            "ContentType": COMMENTS_CONTENT_TYPE,
                        },
                    )
                data = etree.tostring(ct_root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            elif item.filename == "word/_rels/document.xml.rels":
                rels_root = etree.fromstring(data)
                exists = any(
                    el.get("Type") == COMMENTS_REL_TYPE
                    for el in rels_root.findall(f"{{{REL_NS}}}Relationship")
                )
                if not exists:
                    etree.SubElement(
                        rels_root,
                        f"{{{REL_NS}}}Relationship",
                        {
                            "Id": "rIdVendorComments99",
                            "Type": COMMENTS_REL_TYPE,
                            "Target": "comments.xml",
                        },
                    )
                data = etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            zout.writestr(item, data)
        zout.writestr("word/comments.xml", comments_xml_bytes)

    return out_mem.getvalue()


def _add_contract_cover_and_recitals(doc: Document, is_redline: bool) -> None:
    """Add formal aerospace contract cover block, metadata table, and recitals before the Heading 2 clauses."""
    doc.add_heading(
        "CYMBAL AEROSPACE CORPORATION — MASTER AVIONICS PRODUCTION SUBCONTRACT (GP1 / SP3 / SP5)",
        level=1,
    )
    subtitle = (
        "VENDOR REDLINED MARKUP (SUBMITTED BY APEX AEROSYSTEMS INC. LEGAL & GLOBAL TAX)"
        if is_redline
        else "CYMBAL AEROSPACE PROFORMA BASELINE CONTRACT (GP1 REV 2026 + SP3 + SP5)"
    )
    doc.add_paragraph(subtitle)

    cover_rows = [
        ("Master Agreement / Deal ID", "BA-2026-AV-8921"),
        ("Buyer (Prime Airframe OEM)", "Cymbal Aerospace Corporation (Commercial Airplanes Division — Everett & Renton, WA)"),
        ("Seller (Tier-1 Avionics Supplier)", "Apex AeroSystems Inc. (Flight Control Systems Division — Phoenix, AZ)"),
        ("Aircraft Programs Covered", "Cymbal C780 Widebody & C730-MAX Next-Gen Flight Control Electronics (FCC-900)"),
        ("Contract Type & Ceiling Value", "Firm-Fixed-Price (FFP) Master Production Subcontract — $45,000,000.00 USD (300 Shipsets)"),
        ("Governing Proforma Stack", "Cymbal GP1 (General Provisions) + SP3 (International/Tariffs) + SP5 (Cybersecurity/DFARS)"),
        ("Companion SOW & Pricing Schedule", "Cymbal_SOW_Schedule_A_Pricing.docx (SOW-AV-8921 Schedule A — GP1 §3.1 FFP Lock)"),
        ("M365 SharePoint Deal Room", SHAREPOINT_SITE_URL),
    ]
    tbl = doc.add_table(rows=len(cover_rows), cols=2)
    for idx, (k, v) in enumerate(cover_rows):
        tbl.rows[idx].cells[0].text = k
        tbl.rows[idx].cells[1].text = v

    doc.add_heading("RECITALS & CONTRACTUAL BACKGROUND", level=1)
    doc.add_paragraph(
        "WHEREAS, Cymbal Aerospace Corporation (\"Buyer\") designs, manufactures, flight-tests, and delivers "
        "transport-category commercial aircraft certificated under Federal Aviation Administration (\"FAA\") "
        "14 CFR Part 21 Production Certificate No. PC-700; and"
    )
    doc.add_paragraph(
        "WHEREAS, Apex AeroSystems Inc. (\"Seller\") is an accredited Tier-1 aerospace avionics manufacturer "
        "holding AS9100 Rev D quality certification and RTCA DO-178C / DO-254 design assurance capabilities for "
        "Design Assurance Level A (\"DAL-A\") primary flight control electronics; and"
    )
    doc.add_paragraph(
        "WHEREAS, Buyer desires to procure from Seller, and Seller desires to design, qualify, manufacture, test, "
        "and deliver to Buyer, three hundred (300) shipsets of Triple-Redundant Fly-By-Wire Flight Control Computers "
        "(Part No. CYM-FCC-900-A) on a strict Firm-Fixed-Price (\"FFP\") basis pursuant to Statement of Work "
        "SOW-AV-8921 Schedule A ($45,000,000.00 ceiling);"
    )
    doc.add_paragraph(
        "NOW, THEREFORE, in consideration of the mutual covenants and agreements set forth herein, Cymbal Aerospace "
        "Corporation and Apex AeroSystems Inc. hereby agree to the following General Provisions (GP1), International "
        "Special Provisions (SP3), and Cybersecurity Special Provisions (SP5):"
    )


def _add_signature_block(doc: Document) -> None:
    """Append a formal bilateral execution signature block table (no paragraphs after last clause)."""
    sig_table = doc.add_table(rows=4, cols=2)
    sig_table.rows[0].cells[0].text = "BUYER: CYMBAL AEROSPACE CORPORATION"
    sig_table.rows[0].cells[1].text = "SELLER: APEX AEROSYSTEMS INC."
    sig_table.rows[1].cells[0].text = "By: ________________________________________"
    sig_table.rows[1].cells[1].text = "By: ________________________________________"
    sig_table.rows[2].cells[0].text = "Name: Marcus Vance\nTitle: Vice President, Avionics Procurement"
    sig_table.rows[2].cells[1].text = "Name: Victoria Sterling, J.D.\nTitle: SVP & General Counsel, Apex AeroSystems"
    sig_table.rows[3].cells[0].text = "Date: October 2026 | Place: Everett, Washington"
    sig_table.rows[3].cells[1].text = "Date: October 2026 | Place: Phoenix, Arizona"


def build_standard_and_redlined_contracts() -> Tuple[Path, Path]:
    """Generate both Cymbal_GP1_Standard.docx and Vendor_Redlined_GP1.docx (plus compatibility copies)."""
    GOLDEN_DIR.mkdir(parents=True, exist_ok=True)

    # 1. Build Cymbal_GP1_Standard.docx (100% untouched baseline)
    base_doc = Document()
    _add_contract_cover_and_recitals(base_doc, is_redline=False)

    for c in CLAUSES_SPEC:
        if c["part_heading"]:
            base_doc.add_heading(c["part_heading"], level=1)
        h2_text = (
            f"[{c['clause_id']}] {c['clause_ref']} — {c['clause_title']} "
            f"| SME_TEAM={c['sme_team']} | OWNER={c['sme_owner_label']} "
            f"| RISK={c['risk_level']} | REG={c['regulatory_flag']}"
        )
        base_doc.add_heading(h2_text, level=2)
        p = base_doc.add_paragraph()
        baseline_str = "".join(txt for kind, txt in c["segments"] if kind in ("same", "del"))
        _append_run(p._p, baseline_str)

    _add_signature_block(base_doc)

    cymbal_std_path = GOLDEN_DIR / "Cymbal_GP1_Standard.docx"
    boeing_std_path = GOLDEN_DIR / "Boeing_GP1_Standard.docx"
    base_doc.save(str(cymbal_std_path))
    shutil.copy2(cymbal_std_path, boeing_std_path)

    # 2. Build Vendor_Redlined_GP1.docx (with native <w:del>, <w:ins>, and word/comments.xml)
    red_doc = Document()
    _add_contract_cover_and_recitals(red_doc, is_redline=True)

    rev_id = 100
    comment_id = 1
    comments_to_inject: List[Dict[str, Any]] = []

    for c in CLAUSES_SPEC:
        if c["part_heading"]:
            red_doc.add_heading(c["part_heading"], level=1)
        h2_text = (
            f"[{c['clause_id']}] {c['clause_ref']} — {c['clause_title']} "
            f"| SME_TEAM={c['sme_team']} | OWNER={c['sme_owner_label']} "
            f"| RISK={c['risk_level']} | REG={c['regulatory_flag']}"
        )
        red_doc.add_heading(h2_text, level=2)
        p = red_doc.add_paragraph()
        p_el = p._p

        has_comment = bool(c["comment"])
        if has_comment:
            etree.SubElement(
                p_el, f"{{{W_NS}}}commentRangeStart", {f"{{{W_NS}}}id": str(comment_id)}
            )

        for kind, txt in c["segments"]:
            if kind == "same":
                _append_run(p_el, txt)
            elif kind == "del":
                rev_id += 1
                _append_del(p_el, rev_id, txt)
            elif kind == "ins":
                rev_id += 1
                _append_ins(p_el, rev_id, txt)

        if has_comment:
            etree.SubElement(
                p_el, f"{{{W_NS}}}commentRangeEnd", {f"{{{W_NS}}}id": str(comment_id)}
            )
            r_ref = etree.SubElement(p_el, f"{{{W_NS}}}r")
            etree.SubElement(
                r_ref, f"{{{W_NS}}}commentReference", {f"{{{W_NS}}}id": str(comment_id)}
            )
            comments_to_inject.append(
                {
                    "id": comment_id,
                    "author": VENDOR_AUTHOR,
                    "date": REDLINE_DATE,
                    "text": c["comment"],
                }
            )
            comment_id += 1

    _add_signature_block(red_doc)

    buf = BytesIO()
    red_doc.save(buf)
    final_red_bytes = _inject_comments_part(buf.getvalue(), comments_to_inject)

    vendor_red_path = GOLDEN_DIR / "Vendor_Redlined_GP1.docx"
    cymbal_red_path = GOLDEN_DIR / "Cymbal_Vendor_Redlined_GP1.docx"
    vendor_red_path.write_bytes(final_red_bytes)
    shutil.copy2(vendor_red_path, cymbal_red_path)

    return cymbal_std_path, vendor_red_path


def build_companion_sow_schedule_a() -> Path:
    """Generate `Cymbal_SOW_Schedule_A_Pricing.docx` for the SharePoint deal room."""
    doc = Document()
    doc.add_heading(
        "CYMBAL AEROSPACE CORPORATION — STATEMENT OF WORK (SOW-AV-8921) & SCHEDULE A FIRM-FIXED-PRICE TABLE",
        level=1,
    )
    doc.add_paragraph(
        "Master Subcontract ID: BA-2026-AV-8921 | Buyer: Cymbal Aerospace Corporation | "
        "Seller: Apex AeroSystems Inc. | Governing Price Lock: GP1 §3.1 ($45,000,000.00 Ceiling)"
    )

    doc.add_heading("1. SCOPE OF WORK & FLIGHT-CRITICAL DELIVERABLES", level=2)
    doc.add_paragraph(
        "1.1 Seller (Apex AeroSystems Inc.) shall design, qualify, manufacture, environmental-stress-screen (ESS), "
        "and deliver three hundred (300) shipsets of Triple-Redundant Fly-By-Wire Flight Control Computers "
        "(Part Number CYM-FCC-900-A) and embedded RTCA DO-178C Design Assurance Level A (DAL-A) flight control "
        "operational flight program (OFP) software for the Cymbal C780 Widebody and C730-MAX aircraft programs."
    )
    doc.add_paragraph(
        "1.2 Each shipset consists of three (3) independent channel Line Replaceable Units (LRUs) housed in an "
        "ARINC 600 6-MCU chassis utilizing aerospace-grade Ti-6Al-4V titanium alloy enclosures and radiation-hardened "
        "7nm avionics multi-core processors interfacing with Cymbal Aerospace's proprietary ARINC 664 Part 7 / "
        "MIL-STD-1553B flight control data bus."
    )

    doc.add_heading("2. TECHNICAL, QUALITY & REGULATORY FLOWDOWNS", level=2)
    doc.add_paragraph(
        "2.1 Airworthiness & Quality: All units shall be manufactured under Seller's accredited AS9100 Rev D Quality "
        "Management System and Cymbal Quality Specification D6-82479. Pursuant to GP1 §12.2 and FAA 14 CFR Part 21.140, "
        "all cleanroom assembly, burn-in, and Hardware-in-the-Loop (HIL) test stations shall remain open to unannounced "
        "FAA and Cymbal Aerospace source surveillance at all times. Each shipset requires an FAA Form 8130-3 tag."
    )
    doc.add_paragraph(
        "2.2 Cybersecurity & CUI Protection: All flight control source code, bus interface control documents (ICDs), "
        "and qualification test reports constitute Covered Defense Information (CDI) / Controlled Unclassified "
        "Information (CUI) subject to SP5 §3.2, NIST SP 800-171 Rev 2, and mandatory 72-hour cyber incident reporting "
        "under DoD DFARS 252.204-7012."
    )

    doc.add_heading("3. SCHEDULE A — FIRM-FIXED-PRICE (FFP) UNIT PRICING & CEILING LOCK (GP1 §3.1)", level=2)
    doc.add_paragraph(
        "3.1 MANDATORY FIRM-FIXED-PRICE (FFP) LOCK: Pursuant to GP1 §3.1, the unit prices and total Contract ceiling "
        "set forth in Table A-1 below are strictly Firm-Fixed-Price (FFP) Delivered Duty Paid (DDP Incoterms 2020) "
        "Cymbal Aerospace Everett/Renton Final Assembly Facility. Unit prices include all raw materials (including "
        "Ti-6Al-4V titanium forgings and semiconductor wafers), manufacturing, DAL-A software qualification, ATA Spec 300 "
        "Class 1 ESD packaging, freight, insurance, customs clearance, and import duties/tariffs. No Special Provision "
        "(including SP3 §4.1) may impose an uncapped tariff pass-through or surcharge that breaches this $45,000,000.00 "
        "FFP ceiling; any authorized cost-sharing adjustment under SP3 §4.1 must be strictly capped at a maximum "
        "aggregate adjustment of two percent (2.0%) of the SOW Schedule A unit price ($3,000/shipset; $900,000 max)."
    )

    headers = [
        "CLIN",
        "Part Number & Description",
        "Aircraft Program",
        "Delivery Lot Window",
        "Quantity (Shipsets)",
        "FFP Unit Price (DDP)",
        "Extended FFP Ceiling",
    ]
    rows_data = [
        (
            "0001",
            "CYM-FCC-900-A Triple-Redundant Flight Control Computer (Lot 1 LRIP + FAIR)",
            "Cymbal C780 Widebody",
            "Q1 2027 – Q4 2027",
            "75 Shipsets",
            "$150,000.00",
            "$11,250,000.00",
        ),
        (
            "0002",
            "CYM-FCC-900-A Triple-Redundant Flight Control Computer (Lot 2 Full-Rate)",
            "Cymbal C730-MAX",
            "Q1 2028 – Q4 2028",
            "125 Shipsets",
            "$150,000.00",
            "$18,750,000.00",
        ),
        (
            "0003",
            "CYM-FCC-900-A Triple-Redundant Flight Control Computer (Lot 3 Full-Rate)",
            "Cymbal C780 / C730-MAX",
            "Q1 2029 – Q4 2029",
            "100 Shipsets",
            "$150,000.00",
            "$15,000,000.00",
        ),
        (
            "0004",
            "DO-178C DAL-A Software Qualification, AS9102 FAIR & HIL Test Bench Data",
            "All Programs",
            "Included in CLINs 0001–0003",
            "Included",
            "$0.00 (NSP)",
            "$0.00 (Included)",
        ),
        (
            "TOTAL",
            "SOW-AV-8921 SCHEDULE A FIRM-FIXED-PRICE CONTRACT CEILING (GP1 §3.1 LOCK)",
            "C780 / C730-MAX",
            "FY2027 – FY2029",
            "300 Shipsets",
            "$150,000.00 / Shipset",
            "$45,000,000.00",
        ),
    ]

    tbl = doc.add_table(rows=1 + len(rows_data), cols=len(headers))
    for col_i, h in enumerate(headers):
        tbl.rows[0].cells[col_i].text = h
    for row_i, r_vals in enumerate(rows_data, start=1):
        for col_i, val in enumerate(r_vals):
            tbl.rows[row_i].cells[col_i].text = val

    out_path = GOLDEN_DIR / "Cymbal_SOW_Schedule_A_Pricing.docx"
    doc.save(str(out_path))
    return out_path


def build_vendor_redline_transmittal_memo() -> Path:
    """Generate `Apex_AeroSystems_Redline_Transmittal_Memo.docx` for the SharePoint deal room."""
    doc = Document()
    doc.add_heading(
        "APEX AEROSYSTEMS INC. — COMMERCIAL & LEGAL REDLINE TRANSMITTAL MEMORANDUM",
        level=1,
    )
    meta_rows = [
        ("TO:", "Marcus Vance, VP Avionics Procurement — Cymbal Aerospace Corporation"),
        ("FROM:", "Victoria Sterling, J.D., SVP & General Counsel — Apex AeroSystems Inc."),
        ("CC:", "Apex AeroSystems Global Tax, Director of Cleanroom Operations, and Chief Information Security Officer"),
        ("DATE:", "October 15, 2026"),
        ("SUBJECT:", "Vendor Markup & Commercial Exceptions — Deal BA-2026-AV-8921 (FCC-900 Flight Control Computers, $45,000,000)"),
        ("ATTACHMENTS:", "Vendor_Redlined_GP1.docx (with Native Word Track Changes & Margin Comments)"),
    ]
    tbl = doc.add_table(rows=len(meta_rows), cols=2)
    for idx, (k, v) in enumerate(meta_rows):
        tbl.rows[idx].cells[0].text = k
        tbl.rows[idx].cells[1].text = v

    doc.add_heading("1. EXECUTIVE SUMMARY OF APEX AEROSYSTEMS REVIEW", level=2)
    doc.add_paragraph(
        "Apex AeroSystems Inc. is pleased to submit our formal markup of Cymbal Aerospace Corporation's Master Avionics "
        "Production Subcontract (GP1 / SP3 / SP5) for Deal BA-2026-AV-8921 covering 300 shipsets of the CYM-FCC-900-A "
        "Triple-Redundant Flight Control Computer ($45,000,000.00 total program value). To expedite execution prior to "
        "the Q4 Long-Lead Titanium & Wafer Material Release milestone, Apex AeroSystems has ACCEPTED EIGHT (8) OF THE "
        "TWELVE (12) PROFORMA CLAUSES VERBATIM WITHOUT MODIFICATION (GP1 §1.1, GP1 §3.1, GP1 §5.2, GP1 §7.1, GP1 §9.1, "
        "GP1 §15.1, GP1 §28.1, and SP3 §2.1)."
    )
    doc.add_paragraph(
        "Our markup in `Vendor_Redlined_GP1.docx` is strictly limited to four (4) targeted commercial, tax, and "
        "operational adjustments required by our Board of Directors Risk Committee and Security Operations Center:"
    )

    doc.add_heading("2. SUMMARY OF FOUR (4) REQUESTED REDLINE DEVIATIONS", level=2)
    doc.add_paragraph(
        "Exception 1 — GP1 §12.2 (Warranty & FAA Source Inspection): We have modified the hardware/firmware warranty "
        "from 36 months after aircraft delivery to our standard commercial avionics baseline of twelve (12) months from "
        "factory shipment. Additionally, because unannounced walk-in audits disrupt our Class-100 lithography and "
        "automated optical inspection shifts in Phoenix, we request thirty (30) days prior written notice for facility inspections."
    )
    doc.add_paragraph(
        "Exception 2 — GP1 §19 & §24 (Limitation of Liability & IP Indemnity): Uncapped direct, line-stoppage, and "
        "recall liability on a $45M contract exceeds our corporate insurance tower. We have proposed capping aggregate "
        "liability at one hundred percent (100%) of the total Purchase Order value ($45M) and carving out third-party "
        "IP infringement claims arising from the combination of our FCC-900 LRUs with any Cymbal Aerospace or third-party flight control bus."
    )
    doc.add_paragraph(
        "Exception 3 — SP3 §4.1 (Customs, Duties & International Tariffs): Due to acute volatility in Section 232 "
        "tariffs on European aerospace-grade Ti-6Al-4V forgings and Section 301 tariffs on Taiwanese semiconductor "
        "wafers, Apex AeroSystems cannot absorb open-ended tariff spikes under DDP pricing. We have redlined SP3 §4.1 "
        "so that any new or increased Section 232/301 tariffs exceeding five percent (5%) are passed through 100% to Cymbal Aerospace."
    )
    doc.add_paragraph(
        "Exception 4 — SP5 §3.2 (Cybersecurity Incident Reporting — DFARS): While Apex AeroSystems maintains full "
        "NIST SP 800-171 Rev 2 controls, our Security Operations Center (SOC) requires ten (10) business days following "
        "confirmation of a material breach to complete root-cause forensic triage before issuing external notifications, "
        "rather than 72 hours from initial unverified alert discovery."
    )

    doc.add_heading("3. REQUESTED NEXT STEPS", level=2)
    doc.add_paragraph(
        "We welcome a consolidated counter-proposal from Cymbal Aerospace's Procurement, Engineering/FAA, Corporate "
        "Legal, Global Tax, and Cyber Defense teams so we can finalize the Deviation Matrix and execute Deal BA-2026-AV-8921."
    )

    out_path = GOLDEN_DIR / "Apex_AeroSystems_Redline_Transmittal_Memo.docx"
    doc.save(str(out_path))
    return out_path


def update_expected_extracted_diffs(vendor_red_path: Path) -> Path:
    """Run parse_redlined_docx on Vendor_Redlined_GP1.docx and write expected_extracted_diffs.json."""
    extracted = parse_redlined_docx(vendor_red_path)
    assert len(extracted) == 12, f"Expected 12 clauses, got {len(extracted)}"
    redlined_ids = [d.clause_id for d in extracted if d.has_redline]
    assert redlined_ids == ["GP1-12", "GP1-19", "SP3-04", "SP5-03"], f"Unexpected redlined IDs: {redlined_ids}"

    out_json_path = GOLDEN_DIR / "expected_extracted_diffs.json"
    out_json_path.write_text(
        json.dumps([d.model_dump() for d in extracted], indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return out_json_path


def _sql_str(val: str) -> str:
    escaped = val.replace("\\", "\\\\").replace("'", "\\'")
    return f"'{escaped}'"


def update_seed_sql(vendor_red_path: Path) -> Path:
    """Deterministically update sql/02_seed_data.sql with Cymbal Aerospace branding and expanded clauses."""
    extracted = parse_redlined_docx(vendor_red_path)
    diffs_by_id = {d.clause_id: d for d in extracted}
    spec_by_id = {c["clause_id"]: c for c in CLAUSES_SPEC}

    seed_sql_path = REPO_ROOT / "sql" / "02_seed_data.sql"
    existing_sql = seed_sql_path.read_text(encoding="utf-8")

    # Extract Section 2..3 (historical_clause_exceptions + ML.GENERATE_EMBEDDING) from existing SQL
    # and replace Boeing references with Cymbal Aerospace
    sec2_start = existing_sql.index("-- 2. Seed 64 Historical Executed Clause Exceptions")
    sec4_start = existing_sql.index("-- 4. Seed Declarative Cross-Clause Conflict & Regulatory Lockout Rules")
    historical_and_emb_block = existing_sql[sec2_start:sec4_start]
    historical_and_emb_block = (
        historical_and_emb_block.replace("The Boeing Company", "Cymbal Aerospace Corporation")
        .replace("Boeing's", "Cymbal Aerospace's")
        .replace("Boeing", "Cymbal Aerospace")
    )

    # 1. Build Section 1: proforma_clauses
    proforma_rows_sql: List[str] = []
    for c in CLAUSES_SPEC:
        cid = c["clause_id"]
        d = diffs_by_id[cid]
        std_text = d.boeing_baseline_text
        fb1_text = c.get("fallback_tier_1_text") or std_text
        fb2_text = c.get("fallback_tier_2_text") or std_text
        reg_lock = "TRUE" if c["regulatory_lockout"] else "FALSE"
        proforma_rows_sql.append(
            f"  (\n"
            f"    {_sql_str(cid)}, {_sql_str(c['proforma_doc'])}, {_sql_str(c['clause_ref'])}, {_sql_str(c['clause_title'])},\n"
            f"    {_sql_str(c['sme_team'])}, {_sql_str(c['sme_owner_label'])}, {_sql_str(c['risk_level'])},\n"
            f"    {_sql_str(std_text)},\n"
            f"    {_sql_str(fb1_text)},\n"
            f"    {_sql_str(fb2_text)},\n"
            f"    {reg_lock}, {_sql_str(c['regulatory_citation'])},\n"
            f"    {_sql_str(c['guardrail_summary'])},\n"
            f"    []\n"
            f"  )"
        )

    sec1_sql = (
        "-- =============================================================================\n"
        "-- sql/02_seed_data.sql — Seed Canonical Cymbal Aerospace CLM Dataset in Live BigQuery\n"
        "-- Target Dataset: vertex-ai-demos-468803.boeing_clm\n"
        "-- =============================================================================\n\n"
        "-- 1. Seed 12 Canonical Cymbal Aerospace Proforma Clauses (GP1, SP3, SP5)\n"
        "TRUNCATE TABLE `vertex-ai-demos-468803.boeing_clm.proforma_clauses`;\n\n"
        "INSERT INTO `vertex-ai-demos-468803.boeing_clm.proforma_clauses` (\n"
        "  clause_id, proforma_doc, clause_ref, clause_title, sme_team, sme_owner_label,\n"
        "  risk_level, standard_text, fallback_tier_1_text, fallback_tier_2_text,\n"
        "  regulatory_lockout, regulatory_citation, guardrail_summary, embedding\n"
        ")\n"
        "VALUES\n"
        + ",\n".join(proforma_rows_sql)
        + ";\n\n"
    )

    # 4. Build Section 4: clause_conflict_rules
    sp304_harm = spec_by_id["SP3-04"]["fallback_tier_1_text"]
    gp12_std = diffs_by_id["GP1-12"].boeing_baseline_text
    sp503_std = diffs_by_id["SP5-03"].boeing_baseline_text

    sec4_sql = f"""-- 4. Seed Declarative Cross-Clause Conflict & Regulatory Lockout Rules
TRUNCATE TABLE `vertex-ai-demos-468803.boeing_clm.clause_conflict_rules`;

INSERT INTO `vertex-ai-demos-468803.boeing_clm.clause_conflict_rules` (
  rule_id, rule_type, source_clause_id, target_clause_id, source_sme_team, target_sme_team,
  trigger_option_keys, forbidden_patterns, required_patterns, severity,
  conflict_title, conflict_description, harmonized_option_key, harmonized_resolution_text
)
VALUES
  (
    'CONFLICT-SP3-04-GP1-03',
    'CROSS_CLAUSE_CONFLICT',
    'SP3-04',
    'GP1-03',
    'TAX',
    'PROCUREMENT',
    ['accept_vendor'],
    ['exceeding five percent (5%)', 'passed through 100%', 'uncapped'],
    ['capped at a maximum aggregate adjustment of 2.0%'],
    'CRITICAL_CONFLICT',
    'Cross-Clause Conflict: SP3 §4.1 Tariff Pass-Through Violates GP1 §3.1 Firm-Fixed-Price SOW Cap',
    'Global Tax approved a >5% Section 232/301 tariff surcharge pass-through on SP3 §4.1, which directly violates Cymbal Aerospace Procurement Firm-Fixed-Price SOW Cap in GP1 §3.1 ($45M ceiling).',
    'custom_harmonized',
    {_sql_str(sp304_harm)}
  ),
  (
    'LOCKOUT-GP1-12-FAA',
    'REGULATORY_LOCKOUT',
    'GP1-12',
    'GP1-12',
    'ENGINEERING',
    'ENGINEERING',
    ['accept_vendor'],
    ['thirty (30) days prior written notice', 'scheduled inspections upon thirty'],
    ['unannounced source inspections'],
    'REGULATORY_BLOCK',
    'Regulatory Lockout: GP1 §12.2 Prior Notice Violates FAA 14 CFR Part 21.140',
    'FAA 14 CFR Part 21.140 / AS9100 Rev D mandatory unannounced right of access cannot be waived or restricted by 30-day prior written notice.',
    'hold_firm',
    {_sql_str(gp12_std)}
  ),
  (
    'LOCKOUT-SP5-03-DFARS',
    'REGULATORY_LOCKOUT',
    'SP5-03',
    'SP5-03',
    'CYBER',
    'CYBER',
    ['accept_vendor'],
    ['ten (10) business days'],
    ['seventy-two (72) hours'],
    'REGULATORY_BLOCK',
    'Regulatory Lockout: SP5 §3.2 10-Day Delay Violates DoD DFARS 252.204-7012',
    'DoD DFARS 252.204-7012 mandatory 72-hour cyber incident reporting window cannot be extended to 10 business days.',
    'hold_firm',
    {_sql_str(sp503_std)}
  );

"""

    # 5. Build Section 5: active_deals
    sec5_sql = f"""-- 5. Seed Active Deal (BA-2026-AV-8921)
TRUNCATE TABLE `vertex-ai-demos-468803.boeing_clm.active_deals`;

INSERT INTO `vertex-ai-demos-468803.boeing_clm.active_deals` (
  deal_id, vendor_name, supplier_tier, boeing_program, contract_value_usd,
  proforma_stack, pricing_model, sow_cap_clause_ref, sharepoint_folder_url,
  deviation_matrix_xlsx, total_clauses_parsed, untouched_clauses_count,
  redlined_clauses_count, status, updated_at
)
VALUES (
  'BA-2026-AV-8921',
  'Apex AeroSystems Inc.',
  'Tier-1 Avionics',
  'Cymbal C780 / C730-MAX Flight Control Electronics (FCC-900)',
  45000000.0,
  'Cymbal GP1 (General Provisions) + SP3 (International/Tax) + SP5 (Cyber/DFARS)',
  'Firm-Fixed-Price (FFP) — $45,000,000 Ceiling',
  'GP1 §3.1',
  {_sql_str(SHAREPOINT_SITE_URL)},
  'BA-2026-AV-8921_SharePoint_Deviation_Matrix.xlsx',
  12,
  8,
  4,
  'CONFLICT_DETECTED',
  CURRENT_TIMESTAMP()
);

"""

    # 6. Build Section 6: deal_clause_deviations
    initial_sme_meta = {
        "GP1-12": (
            "hold_firm",
            gp12_std,
            "hold_firm",
            gp12_std,
            "APPROVED",
            "Dr. Aris Thorne (Principal FAA DER, Flight Controls QA)",
            "Rejected vendor 12-month warranty and 30-day inspection notice. FAA 14 CFR Part 21.140 and AS9100 Rev D mandate unannounced source inspection access at all times. Precedent shows 88% (14/16) held firm.",
        ),
        "GP1-19": (
            "compromise_tier_1",
            spec_by_id["GP1-19"]["fallback_tier_1_text"],
            "compromise_tier_1",
            spec_by_id["GP1-19"]["fallback_tier_1_text"],
            "APPROVED",
            "Elena Rostova, J.D. (Senior Counsel, Commercial & IP)",
            "Applied Playbook Tier-1 Compromise (69% historical precedent across 11/16 Tier-1 avionics deals): 2x PO ($90M) super-cap on general direct damages while holding uncapped liability for airworthiness recall and IP indemnity, narrowing the IP carveout strictly to Cymbal Aerospace proprietary bus specifications.",
        ),
        "SP3-04": (
            "compromise_tier_1",
            sp304_harm,
            "accept_vendor",
            diffs_by_id["SP3-04"].vendor_redline_text,
            "CONFLICT",
            "Marcus Chen, CPA (Director, Global Tax & Trade)",
            "Accepted vendor >5% Section 232/301 titanium and wafer tariff pass-through to avoid Q3 avionics delivery delay, pending Procurement verification against SOW cap.",
        ),
        "SP5-03": (
            "hold_firm",
            sp503_std,
            "hold_firm",
            sp503_std,
            "APPROVED",
            "Capt. Devon Brooks (Lead, Supplier Cyber Defense)",
            "Held firm on 72-hour incident notification. DFARS 252.204-7012 and NIST SP 800-171 impose a non-negotiable 72-hour DoD reporting window (94% held firm across 15/16 historical contracts).",
        ),
    }

    dev_rows_sql: List[str] = []
    for d in extracted:
        cid = d.clause_id
        if not d.has_redline:
            rec_key = "hold_firm"
            rec_txt = d.boeing_baseline_text
            sel_key = "hold_firm"
            sel_txt = d.boeing_baseline_text
            status = "UNTOUCHED"
            reviewer = "Auto-Verified (No Redline)"
            rationale = (
                "Firm-Fixed-Price SOW Cap ($45M ceiling) locked per GP1 §3.1."
                if cid == "GP1-03"
                else f"Clause matches Cymbal Aerospace Proforma {d.proforma_doc} baseline with zero Track Changes."
            )
        else:
            rec_key, rec_txt, sel_key, sel_txt, status, reviewer, rationale = initial_sme_meta[cid]

        v_comment = d.vendor_comments[0] if d.vendor_comments else ""
        ins_json = json.dumps(d.insertions, ensure_ascii=False)
        del_json = json.dumps(d.deletions, ensure_ascii=False)

        dev_rows_sql.append(
            f"  (\n"
            f"    'BA-2026-AV-8921', {_sql_str(cid)}, {_sql_str(d.proforma_doc)}, {_sql_str(d.clause_ref)}, {_sql_str(d.clause_title)}, {'TRUE' if d.has_redline else 'FALSE'},\n"
            f"    {_sql_str(d.sme_team)}, {_sql_str(d.sme_owner_label)}, {_sql_str(d.risk_level)}, {_sql_str(d.regulatory_flag)},\n"
            f"    {_sql_str(d.boeing_baseline_text)},\n"
            f"    {_sql_str(d.vendor_redline_text)},\n"
            f"    {_sql_str(d.diff_html)},\n"
            f"    {_sql_str(v_comment)}, {_sql_str(ins_json)}, {_sql_str(del_json)},\n"
            f"    {_sql_str(rec_key)},\n"
            f"    {_sql_str(rec_txt)},\n"
            f"    {_sql_str(sel_key)},\n"
            f"    {_sql_str(sel_txt)},\n"
            f"    {_sql_str(status)}, {_sql_str(reviewer)}, {_sql_str(rationale)}, CURRENT_TIMESTAMP()\n"
            f"  )"
        )

    sec6_sql = (
        "-- 6. Seed Deal Clause Deviations (12 Clauses: 8 Untouched, 4 Redlined with Initial Conflict on SP3-04)\n"
        "TRUNCATE TABLE `vertex-ai-demos-468803.boeing_clm.deal_clause_deviations`;\n\n"
        "INSERT INTO `vertex-ai-demos-468803.boeing_clm.deal_clause_deviations` (\n"
        "  deal_id, clause_id, proforma_doc, clause_ref, clause_title, has_redline,\n"
        "  sme_team, sme_owner_label, risk_level, regulatory_flag, boeing_baseline_text,\n"
        "  vendor_redline_text, diff_html, vendor_comment, insertions_json, deletions_json,\n"
        "  recommended_option_key, recommended_counter_text, selected_option_key,\n"
        "  selected_counter_text, sme_status, sme_reviewer, sme_rationale, updated_at\n"
        ")\n"
        "VALUES\n"
        + ",\n".join(dev_rows_sql)
        + ";\n\n"
    )

    # 7. Build Section 7: sme_review_events
    gp19_fb1 = spec_by_id["GP1-19"]["fallback_tier_1_text"]
    sp304_vend = diffs_by_id["SP3-04"].vendor_redline_text

    sec7_sql = f"""-- 7. Seed Immutable SME Review Events Audit Trail
TRUNCATE TABLE `vertex-ai-demos-468803.boeing_clm.sme_review_events`;

INSERT INTO `vertex-ai-demos-468803.boeing_clm.sme_review_events` (
  event_id, deal_id, clause_id, sme_team, actor, action,
  selected_option_key, selected_counter_text, rationale, source_surface, created_at
)
VALUES
  (
    'EVT-8921-001', 'BA-2026-AV-8921', 'ALL', 'PROCUREMENT',
    'Lead Contract Owner (M. Vance)', 'INGESTED_REDLINE', 'hold_firm',
    'Parsed 12 clauses from Vendor_Redlined_GP1.docx: 8 untouched auto-collapsed, 4 redlined routed to Engineering, Legal, Tax, and Cyber SMEs.',
    'Deterministic OOXML Track Changes extraction completed against Cymbal_GP1_Standard.docx.',
    'GEMINI_AGENT_MCP', TIMESTAMP_SUB(CURRENT_TIMESTAMP(), INTERVAL 25 MINUTE)
  ),
  (
    'EVT-8921-002', 'BA-2026-AV-8921', 'GP1-12', 'ENGINEERING',
    'Dr. Aris Thorne (Principal FAA DER, Flight Controls QA)', 'SME_APPROVED', 'hold_firm',
    {_sql_str(gp12_std)},
    'Enforced FAA 14 CFR Part 21.140 unannounced inspection lockout and 36-month warranty (88% held firm across 16 precedents).',
    'EXPERIENCE_LAYER_UI', TIMESTAMP_SUB(CURRENT_TIMESTAMP(), INTERVAL 18 MINUTE)
  ),
  (
    'EVT-8921-003', 'BA-2026-AV-8921', 'GP1-19', 'LEGAL',
    'Elena Rostova, J.D. (Senior Counsel, Commercial & IP)', 'SME_APPROVED', 'compromise_tier_1',
    {_sql_str(gp19_fb1)},
    'Selected Playbook Tier-1 Compromise (2x PO $90M super-cap + narrowed Cymbal Aerospace sole-spec IP carveout; 69% Tier-1 precedent).',
    'EXPERIENCE_LAYER_UI', TIMESTAMP_SUB(CURRENT_TIMESTAMP(), INTERVAL 12 MINUTE)
  ),
  (
    'EVT-8921-004', 'BA-2026-AV-8921', 'SP5-03', 'CYBER',
    'Capt. Devon Brooks (Lead, Supplier Cyber Defense)', 'SME_APPROVED', 'hold_firm',
    {_sql_str(sp503_std)},
    'Enforced DFARS 252.204-7012 mandatory 72-hour cyber incident reporting lockout (94% held firm across 16 precedents).',
    'EXPERIENCE_LAYER_UI', TIMESTAMP_SUB(CURRENT_TIMESTAMP(), INTERVAL 7 MINUTE)
  ),
  (
    'EVT-8921-005', 'BA-2026-AV-8921', 'SP3-04', 'TAX',
    'Marcus Chen, CPA (Director, Global Tax & Trade)', 'CONFLICT_TRIGGERED', 'accept_vendor',
    {_sql_str(sp304_vend)},
    'Global Tax acceptance of uncapped >5% Section 232/301 tariff pass-through triggered CRITICAL_CONFLICT with GP1 §3.1 Firm-Fixed-Price SOW Cap ($45M ceiling).',
    'EXPERIENCE_LAYER_UI', TIMESTAMP_SUB(CURRENT_TIMESTAMP(), INTERVAL 2 MINUTE)
  );
"""

    full_sql = sec1_sql + historical_and_emb_block + sec4_sql + sec5_sql + sec6_sql + sec7_sql
    seed_sql_path.write_text(full_sql, encoding="utf-8")
    return seed_sql_path


def main() -> None:
    cymbal_std, vendor_red = build_standard_and_redlined_contracts()
    sow_path = build_companion_sow_schedule_a()
    memo_path = build_vendor_redline_transmittal_memo()
    diffs_json = update_expected_extracted_diffs(vendor_red)
    seed_sql = update_seed_sql(vendor_red)

    print("Generated Full-Length Cymbal Aerospace Contract Package & Seed SQL:")
    for p in [
        cymbal_std,
        GOLDEN_DIR / "Boeing_GP1_Standard.docx",
        vendor_red,
        GOLDEN_DIR / "Cymbal_Vendor_Redlined_GP1.docx",
        sow_path,
        memo_path,
        diffs_json,
        seed_sql,
    ]:
        print(f"  - {p.name} ({p.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
