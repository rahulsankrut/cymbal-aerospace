"""
specs/04_golden_fixtures/build_golden_fixtures.py
Deterministically builds the executable truth fixtures for Boeing CLM:
1. Boeing_GP1_Standard.docx (12 canonical Boeing clauses across GP1, SP3, SP5)
2. Vendor_Redlined_GP1.docx (12 clauses: 8 untouched, 4 with genuine Word <w:ins>, <w:del>, and <w:comment> XML nodes)
3. expected_extracted_diffs.json (Golden snapshot validated against specs/02_domain_models.py::ExtractedClauseDiff)
"""
from __future__ import annotations

import json
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from lxml import etree

from boeing_clm.models import ExtractedClauseDiff

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
COMMENTS_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
COMMENTS_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"


CANONICAL_CLAUSES: List[Dict[str, Any]] = [
    {
        "clause_id": "GP1-01",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §1.1",
        "clause_title": "Definitions & Order of Precedence",
        "sme_team": "PROCUREMENT",
        "sme_owner_label": "Procurement & Subcontracts",
        "risk_level": "NONE",
        "regulatory_flag": "Boeing Corporate GP1 Standard",
        "has_redline": False,
        "segments": [
            {
                "type": "same",
                "text": "1.1 Buyer means The Boeing Company or its wholly-owned subsidiary identified on the face of the Contract. Seller means the legal entity contracting with Buyer. In the event of any inconsistency among Contract documents, Special Provisions (SP) shall take precedence over General Provisions (GP1), excluding Firm-Fixed-Price caps set forth in GP1 §3.1.",
            }
        ],
        "vendor_comments": [],
    },
    {
        "clause_id": "GP1-03",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §3.1",
        "clause_title": "Firm-Fixed-Price SOW Cap & Pricing Integrity",
        "sme_team": "PROCUREMENT",
        "sme_owner_label": "Procurement & Cost Management",
        "risk_level": "NONE",
        "regulatory_flag": "Firm-Fixed-Price SOW Lock",
        "has_redline": False,
        "segments": [
            {
                "type": "same",
                "text": "3.1 Unless expressly modified by a bilateral Change Order signed by Buyer's Authorized Procurement Representative referencing this GP1 §3.1, all unit prices in Statement of Work (SOW) Schedule A are Firm-Fixed-Price (FFP) and shall not be subject to uncapped surcharges, tariff pass-throughs, or raw material indexing.",
            }
        ],
        "vendor_comments": [],
    },
    {
        "clause_id": "GP1-05",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §5.2",
        "clause_title": "Packing, Marking & Aerospace Preservation",
        "sme_team": "ENGINEERING",
        "sme_owner_label": "Supplier Logistics & Packaging",
        "risk_level": "NONE",
        "regulatory_flag": "ATA Spec 300 / D6-81628",
        "has_redline": False,
        "segments": [
            {
                "type": "same",
                "text": "5.2 Seller shall pack, mark, and ship all Goods in accordance with Boeing Document D6-81628 and ATA Spec 300 electro-static discharge (ESD) protection standards for Class 1 flight-critical avionics.",
            }
        ],
        "vendor_comments": [],
    },
    {
        "clause_id": "GP1-07",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §7.1",
        "clause_title": "Quality Management System (AS9100 Rev D)",
        "sme_team": "ENGINEERING",
        "sme_owner_label": "Supplier Quality Assurance",
        "risk_level": "NONE",
        "regulatory_flag": "AS9100 Rev D / D6-82479",
        "has_redline": False,
        "segments": [
            {
                "type": "same",
                "text": "7.1 Seller shall maintain an accredited AS9100 Rev D Quality Management System and comply with Boeing Quality Management System Requirements for Suppliers (D6-82479) throughout Contract performance.",
            }
        ],
        "vendor_comments": [],
    },
    {
        "clause_id": "GP1-09",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §9.1",
        "clause_title": "Changes & Equitable Adjustment",
        "sme_team": "PROCUREMENT",
        "sme_owner_label": "Procurement & Subcontracts",
        "risk_level": "NONE",
        "regulatory_flag": "Boeing Corporate GP1 Standard",
        "has_redline": False,
        "segments": [
            {
                "type": "same",
                "text": "9.1 Buyer's Authorized Procurement Representative may at any time, by written order, direct changes within the general scope of this Contract in drawings, designs, specifications, or delivery schedules.",
            }
        ],
        "vendor_comments": [],
    },
    {
        "clause_id": "GP1-12",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §12.2",
        "clause_title": "Warranty & FAA Source Inspection",
        "sme_team": "ENGINEERING",
        "sme_owner_label": "Engineering & QA (FAA DER)",
        "risk_level": "HIGH",
        "regulatory_flag": "FAA 14 CFR Part 21 Lockout",
        "has_redline": True,
        "segments": [
            {
                "type": "same",
                "text": "12.2 Seller warrants all Goods against defects in material, workmanship, and avionics firmware for ",
            },
            {"type": "del", "text": "thirty-six (36) months from initial aircraft delivery to Boeing's airline customer"},
            {"type": "ins", "text": "twelve (12) months from factory shipment"},
            {
                "type": "same",
                "text": ". Boeing and FAA representatives may conduct ",
            },
            {"type": "del", "text": "unannounced source inspections at Seller's facility at any time"},
            {"type": "ins", "text": "scheduled inspections upon thirty (30) days prior written notice"},
            {
                "type": "same",
                "text": " to verify AS9100 Rev D and FAA 14 CFR Part 21 production airworthiness compliance.",
            },
        ],
        "vendor_comments": [
            "Apex AeroSystems Legal: Our standard commercial avionics warranty is 12 months from ex-works shipment, and unannounced facility audits disrupt cleanroom assembly shifts."
        ],
    },
    {
        "clause_id": "GP1-15",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §15.1",
        "clause_title": "Termination for Convenience",
        "sme_team": "LEGAL",
        "sme_owner_label": "Corporate Legal & Commercial",
        "risk_level": "NONE",
        "regulatory_flag": "FAR 52.249-2 Aligned",
        "has_redline": False,
        "segments": [
            {
                "type": "same",
                "text": "15.1 Buyer may terminate all or any part of this Contract for convenience by written notice to Seller. Seller shall immediately stop all terminated work and mitigate termination settlement costs.",
            }
        ],
        "vendor_comments": [],
    },
    {
        "clause_id": "GP1-19",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §19 & §24",
        "clause_title": "Limitation of Liability & IP Indemnity",
        "sme_team": "LEGAL",
        "sme_owner_label": "Corporate Legal & IP",
        "risk_level": "CRITICAL",
        "regulatory_flag": "Corporate Legal Playbook Tier-1",
        "has_redline": True,
        "segments": [
            {
                "type": "same",
                "text": "19.1 Seller's liability for direct damages, line-stoppage, and airworthiness recall costs shall ",
            },
            {"type": "del", "text": "be uncapped"},
            {"type": "ins", "text": "not exceed one hundred percent (100%) of the total Purchase Order value ($45M)"},
            {
                "type": "same",
                "text": ". 24.1 Seller shall indemnify Boeing against all third-party intellectual property infringement claims",
            },
            {"type": "del", "text": " including system-level avionics bus integration"},
            {"type": "ins", "text": " except where such claim arises from combination of Goods with any Boeing or third-party flight control bus"},
            {"type": "same", "text": "."},
        ],
        "vendor_comments": [
            "Apex AeroSystems Legal: Board risk policy requires a 1x PO liability cap ($45M) and standard commercial carveout for system-combination IP claims."
        ],
    },
    {
        "clause_id": "GP1-28",
        "proforma_doc": "GP1",
        "clause_ref": "GP1 §28.1",
        "clause_title": "Governing Law & Dispute Resolution",
        "sme_team": "LEGAL",
        "sme_owner_label": "Corporate Legal",
        "risk_level": "NONE",
        "regulatory_flag": "Delaware / Washington Law Standard",
        "has_redline": False,
        "segments": [
            {
                "type": "same",
                "text": "28.1 This Contract and all disputes arising hereunder shall be governed by and construed in accordance with the laws of the State of Delaware, excluding its conflict of laws rules.",
            }
        ],
        "vendor_comments": [],
    },
    {
        "clause_id": "SP3-02",
        "proforma_doc": "SP3",
        "clause_ref": "SP3 §2.1",
        "clause_title": "Export Compliance (ITAR & EAR)",
        "sme_team": "TAX",
        "sme_owner_label": "Global Trade Controls",
        "risk_level": "NONE",
        "regulatory_flag": "22 CFR 120-130 (ITAR) / 15 CFR 730-774 (EAR)",
        "has_redline": False,
        "segments": [
            {
                "type": "same",
                "text": "2.1 Seller shall comply with all applicable U.S. export control laws and regulations, including the International Traffic in Arms Regulations (ITAR) and Export Administration Regulations (EAR).",
            }
        ],
        "vendor_comments": [],
    },
    {
        "clause_id": "SP3-04",
        "proforma_doc": "SP3",
        "clause_ref": "SP3 §4.1",
        "clause_title": "Customs, Duties & International Tariffs",
        "sme_team": "TAX",
        "sme_owner_label": "Global Tax & Trade",
        "risk_level": "MOD",
        "regulatory_flag": "Cross-Clause SOW §3.1 Dependency",
        "has_redline": True,
        "segments": [
            {
                "type": "same",
                "text": "4.1 All unit prices are ",
            },
            {"type": "del", "text": "DDP (Incoterms 2020) Boeing Everett/Renton facility and inclusive of all import duties, customs fees, and tariffs"},
            {"type": "ins", "text": "exclusive of any new or increased Section 232/301 tariffs on titanium forgings or semiconductor wafers exceeding five percent (5%), which shall be passed through 100% to Boeing"},
            {"type": "same", "text": "."},
        ],
        "vendor_comments": [
            "Apex AeroSystems Tax: Due to volatile Section 232/301 duties on European titanium and Taiwanese wafers, Seller requires 100% pass-through of tariff increases above 5%."
        ],
    },
    {
        "clause_id": "SP5-03",
        "proforma_doc": "SP5",
        "clause_ref": "SP5 §3.2",
        "clause_title": "Cybersecurity Incident Reporting (DFARS)",
        "sme_team": "CYBER",
        "sme_owner_label": "Supplier Cyber Defense",
        "risk_level": "HIGH",
        "regulatory_flag": "DFARS 252.204-7012 Lockout",
        "has_redline": True,
        "segments": [
            {
                "type": "same",
                "text": "3.2 Seller shall maintain NIST SP 800-171 controls and report any cybersecurity incident affecting Boeing Covered Defense Information or technical data within ",
            },
            {"type": "del", "text": "seventy-two (72) hours of discovery"},
            {"type": "ins", "text": "ten (10) business days following confirmation of a material breach"},
            {
                "type": "same",
                "text": ", and provide forensic log images upon request.",
            },
        ],
        "vendor_comments": [
            "Apex AeroSystems CISO: Our SOC requires 10 business days to validate root-cause forensics before external customer notification."
        ],
    },
]


def _append_run(p_el: etree._Element, text: str) -> None:
    r = etree.SubElement(p_el, f"{{{W_NS}}}r")
    t = etree.SubElement(r, f"{{{W_NS}}}t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text


def _append_del(p_el: etree._Element, rev_id: int, text: str, author: str = "Apex AeroSystems Legal") -> None:
    del_el = etree.SubElement(
        p_el,
        f"{{{W_NS}}}del",
        {
            f"{{{W_NS}}}id": str(rev_id),
            f"{{{W_NS}}}author": author,
            f"{{{W_NS}}}date": "2026-10-15T14:30:00Z",
        },
    )
    r = etree.SubElement(del_el, f"{{{W_NS}}}r")
    dt = etree.SubElement(r, f"{{{W_NS}}}delText")
    dt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    dt.text = text


def _append_ins(p_el: etree._Element, rev_id: int, text: str, author: str = "Apex AeroSystems Legal") -> None:
    ins_el = etree.SubElement(
        p_el,
        f"{{{W_NS}}}ins",
        {
            f"{{{W_NS}}}id": str(rev_id),
            f"{{{W_NS}}}author": author,
            f"{{{W_NS}}}date": "2026-10-15T14:30:00Z",
        },
    )
    r = etree.SubElement(ins_el, f"{{{W_NS}}}r")
    t = etree.SubElement(r, f"{{{W_NS}}}t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text


def _inject_comments_part(docx_bytes: bytes, comments_list: List[Dict[str, Any]]) -> bytes:
    """Injects word/comments.xml and relationships into a .docx zip archive."""
    in_mem = BytesIO(docx_bytes)
    out_mem = BytesIO()

    nsmap = {"w": W_NS}
    comments_root = etree.Element(f"{{{W_NS}}}comments", nsmap=nsmap)
    for c in comments_list:
        c_el = etree.SubElement(
            comments_root,
            f"{{{W_NS}}}comment",
            {
                f"{{{W_NS}}}id": str(c["id"]),
                f"{{{W_NS}}}author": c["author"],
                f"{{{W_NS}}}date": "2026-10-15T14:35:00Z",
            },
        )
        p_el = etree.SubElement(c_el, f"{{{W_NS}}}p")
        _append_run(p_el, c["text"])

    comments_xml_bytes = etree.tostring(
        comments_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    with zipfile.ZipFile(in_mem, "r") as zin, zipfile.ZipFile(out_mem, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                ct_root = etree.fromstring(data)
                exists = any(
                    el.get("PartName") == "/word/comments.xml" for el in ct_root.findall(f"{{{CT_NS}}}Override")
                )
                if not exists:
                    etree.SubElement(
                        ct_root,
                        f"{{{CT_NS}}}Override",
                        {
                            "PartName": "/word/comments.xml",
                            "ContentType": COMMENTS_CONTENT_TYPE,
                        },
                    )
                data = etree.tostring(ct_root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            elif item.filename == "word/_rels/document.xml.rels":
                rels_root = etree.fromstring(data)
                exists = any(
                    el.get("Type") == COMMENTS_REL_TYPE for el in rels_root.findall(f"{{{REL_NS}}}Relationship")
                )
                if not exists:
                    etree.SubElement(
                        rels_root,
                        f"{{{REL_NS}}}Relationship",
                        {
                            "Id": "rIdComments99",
                            "Type": COMMENTS_REL_TYPE,
                            "Target": "comments.xml",
                        },
                    )
                data = etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            zout.writestr(item, data)
        zout.writestr("word/comments.xml", comments_xml_bytes)

    return out_mem.getvalue()


def build_golden_fixtures(target_dir: Path | None = None) -> List[ExtractedClauseDiff]:
    out_dir = target_dir or Path(__file__).resolve().parent
    out_dir.mkdir(parents=True, exist_ok=True)

    baseline_path = out_dir / "Boeing_GP1_Standard.docx"
    vendor_path = out_dir / "Vendor_Redlined_GP1.docx"
    expected_json_path = out_dir / "expected_extracted_diffs.json"

    # 1. Build Boeing_GP1_Standard.docx
    base_doc = Document()
    base_doc.add_heading("THE BOEING COMPANY — MASTER PRODUCTION CONTRACT (GP1 / SP3 / SP5)", level=1)
    base_doc.add_paragraph(
        "Deal ID: BA-2026-AV-8921 | Supplier: Apex AeroSystems (Tier-1 Avionics) | Program: 787 / 737-MAX Flight Control Display Units | Value: $45,000,000.00 (Firm-Fixed-Price)"
    )

    for c in CANONICAL_CLAUSES:
        meta_header = (
            f"[{c['clause_id']}] {c['clause_ref']} — {c['clause_title']} "
            f"| SME_TEAM={c['sme_team']} | OWNER={c['sme_owner_label']} "
            f"| RISK={c['risk_level']} | REG={c['regulatory_flag']}"
        )
        base_doc.add_heading(meta_header, level=2)
        baseline_text = "".join(
            seg["text"] for seg in c["segments"] if seg["type"] in ("same", "del")
        )
        base_doc.add_paragraph(baseline_text)

    base_doc.save(str(baseline_path))

    # 2. Build Vendor_Redlined_GP1.docx with real <w:del>, <w:ins>, and <w:comment> XML
    red_doc = Document()
    red_doc.add_heading("THE BOEING COMPANY — MASTER PRODUCTION CONTRACT (VENDOR REDLINE SUBMISSION)", level=1)
    red_doc.add_paragraph(
        "Deal ID: BA-2026-AV-8921 | Supplier: Apex AeroSystems (Tier-1 Avionics) | Revision: Vendor Markup v1"
    )

    rev_id = 100
    comment_id = 1
    comments_to_inject: List[Dict[str, Any]] = []
    expected_diffs: List[ExtractedClauseDiff] = []

    for c in CANONICAL_CLAUSES:
        meta_header = (
            f"[{c['clause_id']}] {c['clause_ref']} — {c['clause_title']} "
            f"| SME_TEAM={c['sme_team']} | OWNER={c['sme_owner_label']} "
            f"| RISK={c['risk_level']} | REG={c['regulatory_flag']}"
        )
        red_doc.add_heading(meta_header, level=2)
        p = red_doc.add_paragraph()
        p_el = p._p

        if c["vendor_comments"]:
            etree.SubElement(p_el, f"{{{W_NS}}}commentRangeStart", {f"{{{W_NS}}}id": str(comment_id)})

        insertions: List[str] = []
        deletions: List[str] = []
        diff_parts: List[str] = []
        baseline_parts: List[str] = []
        vendor_parts: List[str] = []

        for seg in c["segments"]:
            if seg["type"] == "same":
                _append_run(p_el, seg["text"])
                diff_parts.append(seg["text"])
                baseline_parts.append(seg["text"])
                vendor_parts.append(seg["text"])
            elif seg["type"] == "del":
                rev_id += 1
                _append_del(p_el, rev_id, seg["text"])
                deletions.append(seg["text"])
                baseline_parts.append(seg["text"])
                diff_parts.append(f'<span class="diff-del">{seg["text"]}</span>')
            elif seg["type"] == "ins":
                rev_id += 1
                _append_ins(p_el, rev_id, seg["text"])
                insertions.append(seg["text"])
                vendor_parts.append(seg["text"])
                diff_parts.append(f'<span class="diff-add">{seg["text"]}</span>')

        if c["vendor_comments"]:
            etree.SubElement(p_el, f"{{{W_NS}}}commentRangeEnd", {f"{{{W_NS}}}id": str(comment_id)})
            r_ref = etree.SubElement(p_el, f"{{{W_NS}}}r")
            rpr_ref = etree.SubElement(r_ref, f"{{{W_NS}}}rPr")
            etree.SubElement(rpr_ref, f"{{{W_NS}}}rStyle", {f"{{{W_NS}}}val": "CommentReference"})
            etree.SubElement(r_ref, f"{{{W_NS}}}commentReference", {f"{{{W_NS}}}id": str(comment_id)})
            comments_to_inject.append(
                {
                    "id": comment_id,
                    "author": "Apex AeroSystems Legal",
                    "text": c["vendor_comments"][0],
                }
            )
            comment_id += 1

        diff_model = ExtractedClauseDiff(
            clause_id=c["clause_id"],
            proforma_doc=c["proforma_doc"],
            clause_ref=c["clause_ref"],
            clause_title=c["clause_title"],
            sme_team=c["sme_team"],
            sme_owner_label=c["sme_owner_label"],
            risk_level=c["risk_level"],
            regulatory_flag=c["regulatory_flag"],
            has_redline=bool(insertions or deletions),
            boeing_baseline_text="".join(baseline_parts),
            vendor_redline_text="".join(vendor_parts),
            insertions=insertions,
            deletions=deletions,
            vendor_comments=list(c["vendor_comments"]),
            diff_html="".join(diff_parts),
        )
        expected_diffs.append(diff_model)

    buf = BytesIO()
    red_doc.save(buf)
    final_docx_bytes = _inject_comments_part(buf.getvalue(), comments_to_inject)
    vendor_path.write_bytes(final_docx_bytes)

    expected_json_path.write_text(
        json.dumps([d.model_dump() for d in expected_diffs], indent=2),
        encoding="utf-8",
    )
    build_portfolio_deal_fixtures(out_dir)
    return expected_diffs


def _build_custom_deal_docx(
    out_path: Path,
    title_heading: str,
    subtitle_line: str,
    clauses_spec: List[Dict[str, Any]],
) -> None:
    doc = Document()
    doc.add_heading(title_heading, level=1)
    doc.add_paragraph(subtitle_line)

    rev_id = 200
    comment_id = 1
    comments_to_inject: List[Dict[str, Any]] = []

    for c in clauses_spec:
        meta_header = (
            f"[{c['clause_id']}] {c['clause_ref']} — {c['clause_title']} "
            f"| SME_TEAM={c['sme_team']} | OWNER={c['sme_owner_label']} "
            f"| RISK={c['risk_level']} | REG={c['regulatory_flag']}"
        )
        doc.add_heading(meta_header, level=2)
        p = doc.add_paragraph()
        p_el = p._p

        comment_items = c.get("comments", [])
        if comment_items:
            etree.SubElement(p_el, f"{{{W_NS}}}commentRangeStart", {f"{{{W_NS}}}id": str(comment_id)})

        for seg in c["segments"]:
            author = seg.get("author", c.get("default_author", "Vendor Legal"))
            if seg["type"] == "same":
                _append_run(p_el, seg["text"])
            elif seg["type"] == "del":
                rev_id += 1
                _append_del(p_el, rev_id, seg["text"], author=author)
            elif seg["type"] == "ins":
                rev_id += 1
                _append_ins(p_el, rev_id, seg["text"], author=author)

        if comment_items:
            etree.SubElement(p_el, f"{{{W_NS}}}commentRangeEnd", {f"{{{W_NS}}}id": str(comment_id)})
            r_ref = etree.SubElement(p_el, f"{{{W_NS}}}r")
            rpr_ref = etree.SubElement(r_ref, f"{{{W_NS}}}rPr")
            etree.SubElement(rpr_ref, f"{{{W_NS}}}rStyle", {f"{{{W_NS}}}val": "CommentReference"})
            etree.SubElement(r_ref, f"{{{W_NS}}}commentReference", {f"{{{W_NS}}}id": str(comment_id)})
            first_c = comment_items[0]
            comments_to_inject.append(
                {
                    "id": comment_id,
                    "author": first_c.get("author", "SME Reviewer"),
                    "text": first_c.get("text", ""),
                }
            )
            comment_id += 1

    buf = BytesIO()
    doc.save(buf)
    final_bytes = _inject_comments_part(buf.getvalue(), comments_to_inject)
    out_path.write_bytes(final_bytes)


def build_portfolio_deal_fixtures(out_dir: Path) -> None:
    """
    Build real golden OOXML .docx fixtures for Deal 2 (BA-2026-TR-5510 — Vanguard Thrust Systems TRA-500)
    and Deal 3 (BA-2026-LG-7740 — Meridian Landing Gear TLG-800) with subsystem-specific text across all 12 clauses.
    """
    tr5510_clauses: List[Dict[str, Any]] = [
        {
            "clause_id": "GP1-01",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §1.1",
            "clause_title": "Definitions & Order of Precedence",
            "sme_team": "PROCUREMENT",
            "sme_owner_label": "Procurement & Subcontracts",
            "risk_level": "NONE",
            "regulatory_flag": "Cymbal Corporate GP1 Standard",
            "segments": [
                {
                    "type": "same",
                    "text": "1.1 Buyer means Cymbal Aerospace Corporation. Seller means Vanguard Thrust Systems Corp. supplying 777X/787 Thrust Reverser Actuators (TRA-500). In the event of any inconsistency among Contract documents, Special Provisions (SP) shall take precedence over General Provisions (GP1), excluding Firm-Fixed-Price caps set forth in GP1 §3.1.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "GP1-03",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §3.1",
            "clause_title": "Firm-Fixed-Price SOW Cap & Pricing Integrity",
            "sme_team": "PROCUREMENT",
            "sme_owner_label": "Procurement Finance (Diego Siciliani)",
            "risk_level": "HIGH",
            "regulatory_flag": "Cross-Department Conflict vs. GP1 §19 Legal Liability Floor",
            "default_author": "Diego Siciliani",
            "segments": [
                {
                    "type": "same",
                    "text": "3.1 All TRA-500 Thrust Reverser Actuator unit prices in Statement of Work (SOW) Schedule A are Firm-Fixed-Price (FFP) totaling $38,000,000, ",
                },
                {
                    "type": "del",
                    "text": "and shall not be subject to uncapped surcharges, tariff pass-throughs, or raw material indexing.",
                },
                {
                    "type": "ins",
                    "text": "less a 3.5% ($1,330,000) commercial volume discount bringing net FFP ceiling to $36,670,000, expressly conditioned upon Buyer accepting Seller's 1.0x ($38,000,000) General Liability Cap in GP1 §19.",
                },
            ],
            "comments": [
                {
                    "author": "Diego Siciliani",
                    "text": "Diego Siciliani (Director, Procurement Finance): Approved 3.5% ($1.33M) commercial unit-price discount on TRA-500 actuators in GP1 §3.1 in exchange for granting Vanguard a 1.0x ($38M) Liability Cap in GP1 §19.",
                }
            ],
        },
        {
            "clause_id": "GP1-05",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §5.2",
            "clause_title": "Packing, Marking & Aerospace Preservation",
            "sme_team": "ENGINEERING",
            "sme_owner_label": "Supplier Logistics & Packaging",
            "risk_level": "NONE",
            "regulatory_flag": "ATA Spec 300 / D6-81628",
            "segments": [
                {
                    "type": "same",
                    "text": "5.2 Seller shall pack, mark, and ship all TRA-500 hydraulic and electromechanical thrust reverser actuator assemblies in accordance with Cymbal Document D6-81628 and ATA Spec 300 corrosion-inhibited hermetically sealed containers.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "GP1-07",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §7.1",
            "clause_title": "Quality Management System (AS9100 Rev D)",
            "sme_team": "ENGINEERING",
            "sme_owner_label": "Supplier Quality Assurance",
            "risk_level": "NONE",
            "regulatory_flag": "AS9100 Rev D / D6-82479",
            "segments": [
                {
                    "type": "same",
                    "text": "7.1 Seller shall maintain an accredited AS9100 Rev D and Nadcap Fluid Distribution Quality Management System for all TRA-500 thrust reverser actuators and comply with D6-82479 throughout Contract performance.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "GP1-09",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §9.1",
            "clause_title": "Changes & Equitable Adjustment",
            "sme_team": "PROCUREMENT",
            "sme_owner_label": "Procurement & Subcontracts",
            "risk_level": "NONE",
            "regulatory_flag": "Cymbal Corporate GP1 Standard",
            "segments": [
                {
                    "type": "same",
                    "text": "9.1 Buyer's Authorized Procurement Representative may at any time, by written order, direct changes within the general scope of this Contract in TRA-500 nacelle interfaces, drawings, or delivery schedules.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "GP1-12",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §12.2",
            "clause_title": "Warranty & FAA Source Inspection",
            "sme_team": "ENGINEERING",
            "sme_owner_label": "Engineering & QA (FAA DER)",
            "risk_level": "HIGH",
            "regulatory_flag": "FAA 14 CFR Part 21 Lockout",
            "default_author": "Vanguard Thrust Legal",
            "segments": [
                {
                    "type": "same",
                    "text": "12.2 Seller warrants all TRA-500 Thrust Reverser Actuator assemblies against defects in hydraulic seals, ballscrews, and workmanship for ",
                },
                {"type": "del", "text": "thirty-six (36) months from initial aircraft delivery to Buyer's airline customer"},
                {"type": "ins", "text": "twelve (12) months from ex-works factory shipment"},
                {
                    "type": "same",
                    "text": ". Buyer and FAA representatives may conduct ",
                },
                {"type": "del", "text": "unannounced source inspections at Seller's facility at any time"},
                {"type": "ins", "text": "scheduled inspections upon thirty (30) days prior written notice"},
                {
                    "type": "same",
                    "text": " to verify AS9100 Rev D and FAA 14 CFR Part 21.140 production airworthiness compliance.",
                },
            ],
            "comments": [
                {
                    "author": "Megan Bowen",
                    "text": "Megan Bowen (Lead Airworthiness & QA DER): Approved Hold Firm — 36-month warranty from aircraft delivery and unannounced FAA 14 CFR Part 21.140 source inspection are non-waivable for TRA-500 thrust reverser hardware.",
                }
            ],
        },
        {
            "clause_id": "GP1-15",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §15.1",
            "clause_title": "Termination for Convenience",
            "sme_team": "LEGAL",
            "sme_owner_label": "Corporate Legal & Commercial",
            "risk_level": "NONE",
            "regulatory_flag": "FAR 52.249-2 Aligned",
            "segments": [
                {
                    "type": "same",
                    "text": "15.1 Buyer may terminate all or any part of this TRA-500 Contract for convenience by written notice to Seller. Seller shall immediately stop all terminated work and mitigate termination settlement costs.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "GP1-19",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §19 & §24",
            "clause_title": "Limitation of Liability & IP Indemnity",
            "sme_team": "LEGAL",
            "sme_owner_label": "Corporate Legal & IP (Miriam Graham)",
            "risk_level": "CRITICAL",
            "regulatory_flag": "Corporate Legal Playbook Minimum 1.5x Liability Cap Guardrail",
            "default_author": "Miriam Graham",
            "segments": [
                {
                    "type": "same",
                    "text": "19.1 Seller's aggregate liability for direct damages, nacelle line-stoppage, and thrust reverser airworthiness recall costs shall ",
                },
                {"type": "del", "text": "not exceed one hundred percent (100% / 1.0x) of the total Purchase Order value ($38,000,000) per GP1 §3.1 commercial discount"},
                {
                    "type": "ins",
                    "text": "not exceed one hundred fifty percent (150% / 1.5x) of total Purchase Order value ($57,000,000), with a 2.5x ($95,000,000) Super-Cap and uncapped carveouts for IP indemnification (§24), flight-safety Airworthiness Directives, and DFARS cybersecurity obligations",
                },
                {"type": "same", "text": "."},
            ],
            "comments": [
                {
                    "author": "Miriam Graham",
                    "text": "Miriam Graham (VP Legal & IP): CONFLICT WITH FINANCE — Legal Playbook strictly requires minimum 1.5x ($57M) Liability Cap on TRA-500 thrust reverser actuators and prohibits commercial cap-for-discount trade-offs in GP1 §3.1. Legal takes precedence.",
                }
            ],
        },
        {
            "clause_id": "GP1-28",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §28.1",
            "clause_title": "Governing Law & Dispute Resolution",
            "sme_team": "LEGAL",
            "sme_owner_label": "Corporate Legal",
            "risk_level": "NONE",
            "regulatory_flag": "Delaware / Washington Law Standard",
            "segments": [
                {
                    "type": "same",
                    "text": "28.1 This TRA-500 Subcontract and all disputes arising hereunder shall be governed by and construed in accordance with the laws of the State of Delaware, excluding its conflict of laws rules.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "SP3-02",
            "proforma_doc": "SP3",
            "clause_ref": "SP3 §2.1",
            "clause_title": "Export Compliance (ITAR & EAR)",
            "sme_team": "TAX",
            "sme_owner_label": "Global Trade Controls",
            "risk_level": "NONE",
            "regulatory_flag": "22 CFR 120-130 (ITAR) / 15 CFR 730-774 (EAR)",
            "segments": [
                {
                    "type": "same",
                    "text": "2.1 Seller shall comply with all applicable U.S. export control laws and regulations (ITAR and EAR Part 774 ECCN 9A991) for all TRA-500 thrust reverser actuation technical data.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "SP3-04",
            "proforma_doc": "SP3",
            "clause_ref": "SP3 §4.1",
            "clause_title": "Customs, Duties & International Tariffs",
            "sme_team": "TAX",
            "sme_owner_label": "Global Tax & Trade",
            "risk_level": "MOD",
            "regulatory_flag": "Cross-Clause SOW §3.1 Dependency",
            "default_author": "Vanguard Thrust Tax",
            "segments": [
                {"type": "same", "text": "4.1 All TRA-500 unit prices are "},
                {
                    "type": "del",
                    "text": "DDP (Incoterms 2020) Cymbal Everett/Renton facility and inclusive of all import duties, customs fees, and tariffs",
                },
                {
                    "type": "ins",
                    "text": "exclusive of any new Section 232/301 tariffs on specialty Inconel/maraging steel actuator forgings exceeding five percent (5%), which shall be shared 50/50 capped at 2.0% ($760,000) of PO value",
                },
                {"type": "same", "text": "."},
            ],
            "comments": [
                {
                    "author": "Pradeep Gupta",
                    "text": "Pradeep Gupta (Director, Global Tax & Trade): Approved Tier-1 50/50 tariff-sharing band above 5.0% threshold on TRA-500 forgings, strictly capped at +2.0% ($760,000) of the $38M FFP ceiling.",
                }
            ],
        },
        {
            "clause_id": "SP5-03",
            "proforma_doc": "SP5",
            "clause_ref": "SP5 §3.2",
            "clause_title": "Cybersecurity Incident Reporting (DFARS)",
            "sme_team": "CYBER",
            "sme_owner_label": "Supplier Cyber Defense",
            "risk_level": "HIGH",
            "regulatory_flag": "DFARS 252.204-7012 Lockout",
            "default_author": "Vanguard Thrust CISO",
            "segments": [
                {
                    "type": "same",
                    "text": "3.2 Seller shall maintain NIST SP 800-171 controls for all TRA-500 digital test benches and report any cybersecurity incident affecting Covered Defense Information within ",
                },
                {"type": "del", "text": "seventy-two (72) hours of discovery"},
                {"type": "ins", "text": "ten (10) business days following confirmation of a material breach"},
                {"type": "same", "text": ", and provide forensic log images upon request."},
            ],
            "comments": [
                {
                    "author": "Alex Wilber",
                    "text": "Alex Wilber (CISO Lead, Supplier Cyber Defense): Approved Hold Firm on mandatory 72-hour DFARS 252.204-7012 cyber incident reporting and 90-day forensic log preservation.",
                }
            ],
        },
    ]

    _build_custom_deal_docx(
        out_path=out_dir / "Vanguard_Working_Review_TR5510.docx",
        title_heading="CYMBAL AEROSPACE — WORKING REVIEW CONTRACT (DEAL BA-2026-TR-5510)",
        subtitle_line="Deal ID: BA-2026-TR-5510 | Supplier: Vanguard Thrust Systems Corp. (Tier-1 Propulsion & Actuation) | Program: 777X / 787 Thrust Reverser Actuators (TRA-500) | Value: $38,000,000.00",
        clauses_spec=tr5510_clauses,
    )

    lg7740_clauses: List[Dict[str, Any]] = [
        {
            "clause_id": "GP1-01",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §1.1",
            "clause_title": "Definitions & Order of Precedence",
            "sme_team": "PROCUREMENT",
            "sme_owner_label": "Procurement & Subcontracts",
            "risk_level": "NONE",
            "regulatory_flag": "Cymbal Corporate GP1 Standard",
            "segments": [
                {
                    "type": "same",
                    "text": "1.1 Buyer means Cymbal Aerospace Corporation. Seller means Meridian Landing Gear Corp. supplying 787-10 Titanium Main Landing Gear Assemblies (TLG-800). In the event of any inconsistency among Contract documents, Special Provisions (SP) shall take precedence over General Provisions (GP1), excluding Firm-Fixed-Price caps set forth in GP1 §3.1.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "GP1-03",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §3.1",
            "clause_title": "Firm-Fixed-Price SOW Cap & Pricing Integrity",
            "sme_team": "PROCUREMENT",
            "sme_owner_label": "Procurement & Cost Management",
            "risk_level": "NONE",
            "regulatory_flag": "Firm-Fixed-Price SOW Lock",
            "segments": [
                {
                    "type": "same",
                    "text": "3.1 Unless expressly modified by a bilateral Change Order signed by Buyer's Authorized Procurement Representative referencing this GP1 §3.1, all TLG-800 Titanium Main Landing Gear unit prices in Statement of Work (SOW) Schedule A are Firm-Fixed-Price ($52,000,000 ceiling) and shall not be subject to uncapped surcharges, tariff pass-throughs, or raw material indexing.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "GP1-05",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §5.2",
            "clause_title": "Packing, Marking & Aerospace Preservation",
            "sme_team": "ENGINEERING",
            "sme_owner_label": "Supplier Logistics & Packaging",
            "risk_level": "NONE",
            "regulatory_flag": "ATA Spec 300 / D6-81628",
            "segments": [
                {
                    "type": "same",
                    "text": "5.2 Seller shall pack, mark, and ship all TLG-800 Ti-10V-2Fe-3Al main landing gear shock struts and truck beams in accordance with Cymbal Document D6-81628 and ATA Spec 300 heavy structural cradle preservation standards.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "GP1-07",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §7.1",
            "clause_title": "Quality Management System (AS9100 Rev D)",
            "sme_team": "ENGINEERING",
            "sme_owner_label": "Supplier Quality Assurance",
            "risk_level": "NONE",
            "regulatory_flag": "AS9100 Rev D / D6-82479",
            "segments": [
                {
                    "type": "same",
                    "text": "7.1 Seller shall maintain an accredited AS9100 Rev D and Nadcap Non-Destructive Testing (Ultrasonic & Eddy Current) Quality Management System for all TLG-800 titanium forgings and comply with D6-82479 throughout Contract performance.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "GP1-09",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §9.1",
            "clause_title": "Changes & Equitable Adjustment",
            "sme_team": "PROCUREMENT",
            "sme_owner_label": "Procurement & Subcontracts",
            "risk_level": "NONE",
            "regulatory_flag": "Cymbal Corporate GP1 Standard",
            "segments": [
                {
                    "type": "same",
                    "text": "9.1 Buyer's Authorized Procurement Representative may at any time, by written order, direct changes within the general scope of this Contract in TLG-800 landing gear trunnion drawings, forging specifications, or delivery schedules.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "GP1-12",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §12.2",
            "clause_title": "Warranty & FAA Source Inspection",
            "sme_team": "ENGINEERING",
            "sme_owner_label": "Engineering & QA (FAA DER)",
            "risk_level": "HIGH",
            "regulatory_flag": "FAA 14 CFR Part 21 Lockout",
            "default_author": "Meridian Landing Gear Legal",
            "segments": [
                {
                    "type": "same",
                    "text": "12.2 Seller warrants all TLG-800 Titanium Main Landing Gear Assemblies against structural fatigue, forging inclusions, and defects in workmanship for ",
                },
                {"type": "del", "text": "thirty-six (36) months from initial aircraft delivery to Buyer's airline customer"},
                {"type": "ins", "text": "eighteen (18) months from ex-works factory shipment"},
                {"type": "same", "text": ". Buyer and FAA representatives may conduct "},
                {"type": "del", "text": "unannounced source inspections at Seller's forge and machining facilities at any time"},
                {"type": "ins", "text": "scheduled inspections upon fifteen (15) business days prior written notice"},
                {"type": "same", "text": " to verify AS9100 Rev D and FAA 14 CFR Part 21.140 primary structural airworthiness compliance."},
            ],
            "comments": [
                {
                    "author": "Megan Bowen",
                    "text": "Megan Bowen (Lead Airworthiness & QA DER): Approved Hold Firm — 36-month warranty from aircraft delivery and unannounced FAA 14 CFR Part 21.140 source inspection are non-waivable for primary TLG-800 titanium landing gear forgings.",
                }
            ],
        },
        {
            "clause_id": "GP1-15",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §15.1",
            "clause_title": "Termination for Convenience",
            "sme_team": "LEGAL",
            "sme_owner_label": "Corporate Legal & Commercial",
            "risk_level": "NONE",
            "regulatory_flag": "FAR 52.249-2 Aligned",
            "segments": [
                {
                    "type": "same",
                    "text": "15.1 Buyer may terminate all or any part of this TLG-800 Contract for convenience by written notice to Seller. Seller shall immediately stop all terminated work and mitigate termination settlement costs.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "GP1-19",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §19 & §24",
            "clause_title": "Limitation of Liability & IP Indemnity",
            "sme_team": "LEGAL",
            "sme_owner_label": "Corporate Legal & IP",
            "risk_level": "CRITICAL",
            "regulatory_flag": "Corporate Legal Playbook Tier-1",
            "default_author": "Meridian Landing Gear Legal",
            "segments": [
                {
                    "type": "same",
                    "text": "19.1 Seller's liability for direct damages, final assembly line-stoppage, and TLG-800 landing gear Airworthiness Directive recall costs shall ",
                },
                {"type": "del", "text": "be uncapped"},
                {"type": "ins", "text": "not exceed one hundred percent (100%) of the total Purchase Order value ($52,000,000)"},
                {
                    "type": "same",
                    "text": ". 24.1 Seller shall indemnify Buyer against all third-party intellectual property infringement claims",
                },
                {"type": "del", "text": " including wheel-well kinematic retraction integration"},
                {"type": "ins", "text": " except where such claim arises from combination of TLG-800 assemblies with Buyer airframe hydraulic systems"},
                {"type": "same", "text": "."},
            ],
            "comments": [
                {
                    "author": "Miriam Graham",
                    "text": "Miriam Graham (VP Legal & IP): Recommended Tier-1 Super-Cap (2.0x / $104,000,000 General Cap with 2.5x / $130,000,000 Super-Cap) with mandatory uncapped carveouts for IP indemnity and flight-safety AD recalls — ready for Contract Owner final sign-off.",
                }
            ],
        },
        {
            "clause_id": "GP1-28",
            "proforma_doc": "GP1",
            "clause_ref": "GP1 §28.1",
            "clause_title": "Governing Law & Dispute Resolution",
            "sme_team": "LEGAL",
            "sme_owner_label": "Corporate Legal",
            "risk_level": "NONE",
            "regulatory_flag": "Delaware / Washington Law Standard",
            "segments": [
                {
                    "type": "same",
                    "text": "28.1 This TLG-800 Subcontract and all disputes arising hereunder shall be governed by and construed in accordance with the laws of the State of Delaware, excluding its conflict of laws rules.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "SP3-02",
            "proforma_doc": "SP3",
            "clause_ref": "SP3 §2.1",
            "clause_title": "Export Compliance (ITAR & EAR)",
            "sme_team": "TAX",
            "sme_owner_label": "Global Trade Controls",
            "risk_level": "NONE",
            "regulatory_flag": "22 CFR 120-130 (ITAR) / 15 CFR 730-774 (EAR)",
            "segments": [
                {
                    "type": "same",
                    "text": "2.1 Seller shall comply with all applicable U.S. export control laws and regulations (ITAR and EAR) for all TLG-800 landing gear metallurgy and machining technical data.",
                }
            ],
            "comments": [],
        },
        {
            "clause_id": "SP3-04",
            "proforma_doc": "SP3",
            "clause_ref": "SP3 §4.1",
            "clause_title": "Customs, Duties & International Tariffs",
            "sme_team": "TAX",
            "sme_owner_label": "Global Tax & Trade",
            "risk_level": "MOD",
            "regulatory_flag": "Cross-Clause SOW §3.1 Dependency",
            "default_author": "Meridian Landing Gear Tax",
            "segments": [
                {"type": "same", "text": "4.1 All TLG-800 unit prices are "},
                {
                    "type": "del",
                    "text": "DDP (Incoterms 2020) Cymbal Everett/Charleston facility and inclusive of all import duties, customs fees, and tariffs",
                },
                {
                    "type": "ins",
                    "text": "exclusive of any new or increased Section 232/301 tariffs on VSMPO/European titanium billet forgings exceeding five percent (5%), which shall be passed through 100% to Buyer",
                },
                {"type": "same", "text": "."},
            ],
            "comments": [
                {
                    "author": "Pradeep Gupta",
                    "text": "Pradeep Gupta (Director, Global Tax & Trade): Approved Tier-1 50/50 tariff-sharing band above 5.0% threshold on TLG-800 titanium forgings, strictly capped at +2.0% ($1,040,000) of the $52M Firm-Fixed-Price SOW.",
                }
            ],
        },
        {
            "clause_id": "SP5-03",
            "proforma_doc": "SP5",
            "clause_ref": "SP5 §3.2",
            "clause_title": "Cybersecurity Incident Reporting (DFARS)",
            "sme_team": "CYBER",
            "sme_owner_label": "Supplier Cyber Defense",
            "risk_level": "HIGH",
            "regulatory_flag": "DFARS 252.204-7012 Lockout",
            "default_author": "Meridian Landing Gear CISO",
            "segments": [
                {
                    "type": "same",
                    "text": "3.2 Seller shall maintain NIST SP 800-171 controls for all CNC 5-axis machining cells and report any cybersecurity incident affecting Covered Defense Information or TLG-800 technical data within ",
                },
                {"type": "del", "text": "seventy-two (72) hours of discovery"},
                {"type": "ins", "text": "seven (7) business days following confirmation of a material breach"},
                {"type": "same", "text": ", and provide forensic log images upon request."},
            ],
            "comments": [
                {
                    "author": "Alex Wilber",
                    "text": "Alex Wilber (CISO Lead, Supplier Cyber Defense): Approved Hold Firm on mandatory 72-hour DFARS 252.204-7012 cyber incident reporting and 90-day raw forensic image preservation.",
                }
            ],
        },
    ]

    _build_custom_deal_docx(
        out_path=out_dir / "Meridian_Working_Review_LG7740.docx",
        title_heading="CYMBAL AEROSPACE — WORKING REVIEW CONTRACT (DEAL BA-2026-LG-7740)",
        subtitle_line="Deal ID: BA-2026-LG-7740 | Supplier: Meridian Landing Gear Corp. (Tier-1 Structural & Landing Gear) | Program: 787-10 Titanium Main Landing Gear Assemblies (TLG-800) | Value: $52,000,000.00",
        clauses_spec=lg7740_clauses,
    )


if __name__ == "__main__":
    diffs = build_golden_fixtures()
    redlined = [d.clause_id for d in diffs if d.has_redline]
    untouched = [d.clause_id for d in diffs if not d.has_redline]
    print(f"Generated {len(diffs)} clauses: {len(untouched)} untouched {untouched}, {len(redlined)} redlined {redlined}")
